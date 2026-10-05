#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_bmc_splits.py
===================

R2 - LEAKAGE-SAFE PARTITIONING

Partitions the eight federated clients into train, validation and test. The unit of
assignment is the GROUP, never the image: whatever key R1 fixed for a client -
patient, case, or the image itself where the release provides no identifier - is
assigned whole, so no group can appear in two partitions. This replaces the random
image-level split used previously.

    GROUP-AWARE   groups are assigned, images follow their group
    CLASS-AWARE   groups are stratified by their own label profile (all benign,
                  mixed, all malignant) and then balanced so the class ratio of each
                  partition tracks the client's overall ratio
    DETERMINISTIC one fixed seed; the assignment is a pure function of the cohort
                  manifest and the seed, and the script proves this by recomputing it

WHY STRATIFY BY GROUP PROFILE AND NOT BY IMAGE LABEL
    A patient can carry benign findings on one breast and malignant on the other:
    42 patients in CBIS-DDSM, 3 in INbreast, 1 in AIISLAB, and among the external
    sets 34 in CDD-CESM and 5 in BUS-UCLM. Such a group cannot be placed in a class.
    Stratifying on the group's own profile keeps those groups intact and still lets
    the class balance be controlled, because the balancing step tracks the malignant
    count each assignment carries rather than assuming one label per group.

THE EXTERNAL SETS ARE NOT PARTITIONED
    BLUIS, BUS-UCLM and CDD-CESM receive the single assignment EXTERNAL_TEST_ONLY.
    They contribute no training data, and they must not be consulted when selecting a
    round, a checkpoint, a threshold or any hyperparameter. The manifest records that
    constraint on every external record and the validation below asserts that none of
    them carries a train or validation assignment.

WHAT THIS SCRIPT PROVES BEFORE IT WRITES
    for every client, the three pairwise group-key intersections are empty
    every record is assigned exactly once, and no image path appears twice
    no external record is assigned to train or validation
    the assignment is reproducible from the seed alone

Usage:
    python build_bmc_splits.py
    python build_bmc_splits.py --ratios 0.70 0.15 0.15 --seed 20260929
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import sys
import traceback
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

DEFAULT_PREP_ROOT = (
    r"D:\PhD\002 FL Mul BC\01 ready\Springer\BMC Med Imaging"
    r"\Revision 1\dataset preparing"
)
BANK_NAME = "CANONICAL_DATASET_BANK"
COHORT_DIR = "_bmc_cohort"
OUT_DIRNAME = "_bmc_splits"
COHORT_FILE = "BMC_COHORT_MANIFEST.csv"
SPLITS = ("train", "validation", "test")
EXTERNAL_ASSIGNMENT = "EXTERNAL_TEST_ONLY"
EXTERNAL_POLICY = ("held out entirely; contributes no training data and must not be "
                   "used to select a round, checkpoint, threshold or any "
                   "hyperparameter")
DEFAULT_SEED = 20260929

SPLIT_FIELDS = [
    "client_id", "client_name", "dataset", "modality", "paper_dataset_id", "role",
    "split", "split_unit", "usage_policy",
    "canonical_record_id", "sample_id",
    "canonical_image_path", "canonical_mask_path",
    "canonical_image_abspath", "canonical_mask_abspath",
    "label", "label_binary",
    "group_key", "group_unit", "group_size", "group_malignant_fraction",
    "group_label_profile", "group_provenance",
    "source_record_id", "seed", "assignment_rule",
]
SUMMARY_FIELDS = [
    "paper_dataset_id", "client_id", "dataset", "modality", "role", "split_unit",
    "train_images", "train_benign", "train_malignant", "train_groups",
    "val_images", "val_benign", "val_malignant", "val_groups",
    "test_images", "test_benign", "test_malignant", "test_groups",
    "total_images", "total_groups",
    "train_pct", "val_pct", "test_pct",
    "overall_malignant_pct", "train_malignant_pct", "val_malignant_pct",
    "test_malignant_pct", "max_class_ratio_deviation_pp", "usage_policy",
]


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def sha256_text(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def read_rows(p: Path) -> List[dict]:
    with p.open(encoding="utf-8-sig", errors="replace") as fh:
        s = fh.read(8192)
        fh.seek(0)
        if not s.strip():
            return []
        d = ";" if s.count(";") > s.count(",") else ","
        return list(csv.DictReader(fh, delimiter=d))


def write_csv(path: Path, fields: Sequence[str], rows: Sequence[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(fields), extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)
    print(f"WROTE {path.name:38s} ({len(rows)} data rows)")


def write_xlsx(path: Path, fields: Sequence[str], rows: Sequence[dict],
               wide: Sequence[str] = ()) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print(f"      openpyxl missing; {path.name} skipped")
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = path.stem[:31]
    ws.append(list(fields))
    for r in rows:
        ws.append([r.get(h, "") for h in fields])
    for c in range(1, len(fields) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5496")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = (
            52 if fields[c - 1] in wide else 15)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"WROTE {path.name}")


def profile_of(n_mal: int, n: int) -> str:
    if n_mal == 0:
        return "ALL_BENIGN"
    if n_mal == n:
        return "ALL_MALIGNANT"
    return "MIXED"


def assign_groups(groups: List[Tuple[str, int, int]], ratios: Tuple[float, ...],
                  seed: int, client_key: str) -> Dict[str, str]:
    """Assign whole groups to train/validation/test.

    Deterministic, and deterministic per client: the RNG is seeded with the run seed
    combined with the client key, so adding or removing a client cannot change the
    partition of any other client.

    Within each label profile the groups are shuffled, then the largest are placed
    first. Each group goes to the split whose image quota is furthest from being met;
    ties are broken by whichever split is furthest below its share of malignant
    images. Placing large groups first matters because one CBIS-DDSM patient carries
    up to 12 images and one BUS-UCLM patient 25, and a late large group would
    otherwise overshoot a small quota by a wide margin.
    """
    out: Dict[str, str] = {}
    rng = random.Random(f"{seed}::{client_key}")
    by_profile: Dict[str, List[Tuple[str, int, int]]] = defaultdict(list)
    for k, n, m in groups:
        by_profile[profile_of(m, n)].append((k, n, m))
    # Each stratum is distributed across the three splits on its OWN quota. A single
    # shared image quota does not work: filling train first would send every
    # all-malignant group there and leave validation and test purely benign, which is
    # exactly what a naive greedy assignment produces here. Per-stratum quotas make
    # each split receive its share of benign, mixed and malignant groups alike, so
    # the class ratio follows the client's ratio by construction.
    for prof in ("MIXED", "ALL_MALIGNANT", "ALL_BENIGN"):
        bucket = sorted(by_profile.get(prof, []), key=lambda g: (g[0],))
        if not bucket:
            continue
        rng.shuffle(bucket)
        bucket.sort(key=lambda g: -g[1])
        p_img = sum(n for _k, n, _m in bucket)
        p_mal = sum(m for _k, _n, m in bucket)
        want_img = {s: p_img * r for s, r in zip(SPLITS, ratios)}
        want_mal = {s: p_mal * r for s, r in zip(SPLITS, ratios)}
        have_img: Dict[str, int] = {s: 0 for s in SPLITS}
        have_mal: Dict[str, int] = {s: 0 for s in SPLITS}
        for k, n, m in bucket:
            def deficit(s: str) -> Tuple[float, float, str]:
                return (-(want_img[s] - have_img[s]),
                        -(want_mal[s] - have_mal[s]), s)
            target = min(SPLITS, key=deficit)
            out[k] = target
            have_img[target] += n
            have_mal[target] += m
    return out


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="R2 - leakage-safe partitioning.")
    ap.add_argument("--prep-root", default=DEFAULT_PREP_ROOT)
    ap.add_argument("--ratios", nargs=3, type=float, default=[0.70, 0.15, 0.15],
                    metavar=("TRAIN", "VAL", "TEST"))
    ap.add_argument("--seed", type=int, default=DEFAULT_SEED)
    a = ap.parse_args(argv)
    prep = Path(a.prep_root)
    bank = prep / BANK_NAME
    cohort_p = bank / COHORT_DIR / COHORT_FILE
    out = bank / OUT_DIRNAME
    ratios = tuple(a.ratios)

    print("=" * 100)
    print("R2 - LEAKAGE-SAFE PARTITIONING: group-aware, class-aware, fixed seed")
    print("=" * 100)
    print(f"generated : {now_iso()}")
    print(f"cohort    : {cohort_p}")
    print(f"ratios    : train {ratios[0]:.2f} / validation {ratios[1]:.2f} / "
          f"test {ratios[2]:.2f}   seed {a.seed}")
    print("unit      : the GROUP fixed by R1, assigned whole; images never split")
    if not cohort_p.is_file():
        print(f"BLOCKED: run R1 first; {COHORT_FILE} not found")
        return 2
    if abs(sum(ratios) - 1.0) > 1e-9:
        print(f"BLOCKED: the ratios sum to {sum(ratios)}, not 1.0")
        return 2

    cohort = read_rows(cohort_p)
    clients = [r for r in cohort if r.get("role") == "CLIENT"]
    external = [r for r in cohort if r.get("role") == "EXTERNAL"]
    print(f"cohort    : {len(cohort)} record(s) = {len(clients)} client + "
          f"{len(external)} external")

    records: List[dict] = []
    summary: List[dict] = []
    problems: List[str] = []
    by_ds: Dict[str, List[dict]] = defaultdict(list)
    for r in clients:
        by_ds[r["dataset"]].append(r)
    order = sorted(by_ds, key=lambda d: int(by_ds[d][0]["client_id"]))

    for ds in order:
        rows = by_ds[ds]
        cid = rows[0]["client_id"]
        gmap: Dict[str, List[dict]] = defaultdict(list)
        for r in rows:
            gmap[r["group_key"]].append(r)
        groups = [(k, len(v), sum(1 for x in v if x["label"] == "MALIGNANT"))
                  for k, v in gmap.items()]
        placement = assign_groups(groups, ratios, a.seed, ds)
        gsize = {k: n for k, n, _m in groups}
        gmal = {k: m for k, _n, m in groups}
        for k, v in gmap.items():
            sp = placement[k]
            for r in v:
                records.append({**{f: r.get(f, "") for f in
                                   ("client_id", "client_name", "dataset", "modality",
                                    "paper_dataset_id", "role",
                                    "canonical_record_id", "sample_id",
                                    "canonical_image_path", "canonical_mask_path",
                                    "canonical_image_abspath",
                                    "canonical_mask_abspath", "label", "label_binary",
                                    "group_key", "group_unit", "group_provenance",
                                    "source_record_id")},
                                "split": sp, "split_unit": r.get("group_unit", ""),
                                "usage_policy": "federated client",
                                "group_size": gsize[k],
                                "group_malignant_fraction": round(
                                    gmal[k] / gsize[k], 4),
                                "group_label_profile": profile_of(gmal[k], gsize[k]),
                                "seed": a.seed,
                                "assignment_rule": (
                                    "GROUP_AWARE_CLASS_AWARE_DETERMINISTIC_V1")})
        stat: Dict[str, Dict[str, int]] = {
            s: {"img": 0, "b": 0, "m": 0, "grp": 0} for s in SPLITS}
        for k, v in gmap.items():
            s = placement[k]
            stat[s]["grp"] += 1
            stat[s]["img"] += len(v)
            stat[s]["b"] += sum(1 for x in v if x["label"] == "BENIGN")
            stat[s]["m"] += sum(1 for x in v if x["label"] == "MALIGNANT")
        tot = len(rows)
        omal = 100.0 * sum(1 for r in rows if r["label"] == "MALIGNANT") / tot
        devs = []
        for s in SPLITS:
            if stat[s]["img"]:
                devs.append(abs(100.0 * stat[s]["m"] / stat[s]["img"] - omal))
        summary.append({
            "paper_dataset_id": rows[0].get("paper_dataset_id", ""),
            "client_id": cid, "dataset": ds, "modality": rows[0]["modality"],
            "role": "CLIENT", "split_unit": rows[0].get("group_unit", ""),
            "train_images": stat["train"]["img"], "train_benign": stat["train"]["b"],
            "train_malignant": stat["train"]["m"], "train_groups": stat["train"]["grp"],
            "val_images": stat["validation"]["img"],
            "val_benign": stat["validation"]["b"],
            "val_malignant": stat["validation"]["m"],
            "val_groups": stat["validation"]["grp"],
            "test_images": stat["test"]["img"], "test_benign": stat["test"]["b"],
            "test_malignant": stat["test"]["m"], "test_groups": stat["test"]["grp"],
            "total_images": tot, "total_groups": len(gmap),
            "train_pct": round(100.0 * stat["train"]["img"] / tot, 1),
            "val_pct": round(100.0 * stat["validation"]["img"] / tot, 1),
            "test_pct": round(100.0 * stat["test"]["img"] / tot, 1),
            "overall_malignant_pct": round(omal, 1),
            "train_malignant_pct": round(
                100.0 * stat["train"]["m"] / stat["train"]["img"], 1)
            if stat["train"]["img"] else 0,
            "val_malignant_pct": round(
                100.0 * stat["validation"]["m"] / stat["validation"]["img"], 1)
            if stat["validation"]["img"] else 0,
            "test_malignant_pct": round(
                100.0 * stat["test"]["m"] / stat["test"]["img"], 1)
            if stat["test"]["img"] else 0,
            "max_class_ratio_deviation_pp": round(max(devs), 1) if devs else 0,
            "usage_policy": "federated client",
        })
        print(f"\n  Client {cid}  {ds:11s} {tot:>5d} image(s) in {len(gmap):>5d} "
              f"{rows[0].get('group_unit', ''):8s} group(s)")
        for s in SPLITS:
            d = stat[s]
            print(f"      {s:11s} {d['img']:>5d} image(s) "
                  f"({100.0 * d['img'] / tot:>5.1f}%)  {d['b']:>5d} B / {d['m']:>5d} M"
                  f"   {d['grp']:>5d} group(s)")
        if min(stat[s]["img"] for s in SPLITS) == 0:
            problems.append(f"{ds}: a partition is empty")
        for s in SPLITS:
            if stat[s]["img"] and (stat[s]["b"] == 0 or stat[s]["m"] == 0):
                problems.append(f"{ds}: the {s} partition holds only one class")

    for r in external:
        records.append({**{f: r.get(f, "") for f in
                           ("client_id", "client_name", "dataset", "modality",
                            "paper_dataset_id", "role", "canonical_record_id",
                            "sample_id", "canonical_image_path",
                            "canonical_mask_path", "canonical_image_abspath",
                            "canonical_mask_abspath", "label", "label_binary",
                            "group_key", "group_unit", "group_provenance",
                            "source_record_id")},
                        "split": EXTERNAL_ASSIGNMENT,
                        "split_unit": r.get("group_unit", ""),
                        "usage_policy": EXTERNAL_POLICY,
                        "group_size": "", "group_malignant_fraction": "",
                        "group_label_profile": "", "seed": a.seed,
                        "assignment_rule": "NOT_PARTITIONED_EXTERNAL_HELD_OUT"})
    for ds in sorted({r["dataset"] for r in external}):
        rows = [r for r in external if r["dataset"] == ds]
        summary.append({
            "paper_dataset_id": "", "client_id": "", "dataset": ds,
            "modality": rows[0]["modality"], "role": "EXTERNAL",
            "split_unit": rows[0].get("group_unit", ""),
            "train_images": 0, "train_benign": 0, "train_malignant": 0,
            "train_groups": 0, "val_images": 0, "val_benign": 0, "val_malignant": 0,
            "val_groups": 0,
            "test_images": len(rows),
            "test_benign": sum(1 for r in rows if r["label"] == "BENIGN"),
            "test_malignant": sum(1 for r in rows if r["label"] == "MALIGNANT"),
            "test_groups": len({r["group_key"] for r in rows}),
            "total_images": len(rows),
            "total_groups": len({r["group_key"] for r in rows}),
            "train_pct": 0, "val_pct": 0, "test_pct": 100.0,
            "overall_malignant_pct": round(
                100.0 * sum(1 for r in rows if r["label"] == "MALIGNANT")
                / len(rows), 1),
            "train_malignant_pct": 0, "val_malignant_pct": 0,
            "test_malignant_pct": round(
                100.0 * sum(1 for r in rows if r["label"] == "MALIGNANT")
                / len(rows), 1),
            "max_class_ratio_deviation_pp": 0, "usage_policy": EXTERNAL_POLICY})

    # ---- the disjointness proof, stated pair by pair ---------------------------
    print("\n" + "=" * 100)
    print("GROUP-KEY DISJOINTNESS, verified per client")
    print("=" * 100)
    leaks: List[str] = []
    for ds in order:
        keys = {s: {r["group_key"] for r in records
                    if r["dataset"] == ds and r["split"] == s} for s in SPLITS}
        pairs = [("train", "validation"), ("train", "test"),
                 ("validation", "test")]
        line = []
        for x, y in pairs:
            inter = keys[x] & keys[y]
            line.append(f"{x[:5]}\u2229{y[:4]}={len(inter)}")
            if inter:
                leaks.append(f"{ds}: {x} and {y} share {len(inter)} group(s): "
                             f"{sorted(inter)[:3]}")
        print(f"  {ds:12s} " + "   ".join(line)
              + f"    all three empty: {'YES' if not any(keys[x] & keys[y] for x, y in pairs) else 'NO'}")

    # ---- reproducibility: recompute and compare --------------------------------
    same = True
    for ds in order:
        rows = by_ds[ds]
        gmap2: Dict[str, List[dict]] = defaultdict(list)
        for r in rows:
            gmap2[r["group_key"]].append(r)
        g2 = [(k, len(v), sum(1 for x in v if x["label"] == "MALIGNANT"))
              for k, v in gmap2.items()]
        p2 = assign_groups(g2, ratios, a.seed, ds)
        first = {r["group_key"]: r["split"] for r in records
                 if r["dataset"] == ds}
        if p2 != first:
            same = False
            problems.append(f"{ds}: recomputing the assignment gave a different result")

    seen_once = Counter((r["dataset"], r["canonical_record_id"])
                        for r in records)
    dup_rec = [k for k, v in seen_once.items() if v > 1]
    img_paths = Counter(r["canonical_image_abspath"] for r in records
                        if r["canonical_image_abspath"])
    ext_bad = [r["dataset"] for r in records
               if r["role"] == "EXTERNAL" and r["split"] != EXTERNAL_ASSIGNMENT]
    cli_bad = [r["dataset"] for r in records
               if r["role"] == "CLIENT" and r["split"] not in SPLITS]

    checks = [
        ("every group key is disjoint across train, validation and test, for every "
         "client", not leaks),
        ("every cohort record received exactly one assignment",
         len(records) == len(cohort) and not dup_rec),
        ("no canonical image is assigned twice",
         not [k for k, v in img_paths.items() if v > 1]),
        ("every client record is in train, validation or test", not cli_bad),
        (f"every external record is {EXTERNAL_ASSIGNMENT} and none is in train or "
         "validation", not ext_bad),
        ("no external record contributes to training or to model selection",
         not [r for r in records if r["role"] == "EXTERNAL"
              and r["split"] in ("train", "validation")]),
        ("the assignment is reproducible from the seed alone", same),
        ("no partition is empty and none holds a single class", not problems),
    ]
    print("\n" + "=" * 100)
    for d, ok in checks:
        print(f"  [{'PASS' if ok else 'FAIL'}]  {d}")
    for x in (leaks + problems)[:8]:
        print(f"     {x}")
    failed = [d for d, ok in checks if not ok]

    cl = [s for s in summary if s["role"] == "CLIENT"]
    print("\n" + "=" * 100)
    print("PARTITION SUMMARY - eight federated clients")
    print("=" * 100)
    print(f"{'ID':4s} {'dataset':11s} {'unit':8s} {'train':>14s} {'val':>13s} "
          f"{'test':>13s} {'M% tr/va/te':>18s} {'dev':>5s}")
    print("-" * 100)
    for s in cl:
        print(f"{str(s['paper_dataset_id']):4s} {s['dataset']:11s} "
              f"{s['split_unit']:8s} "
              f"{s['train_images']:>5d}({s['train_groups']:>4d}g) "
              f"{s['val_images']:>5d}({s['val_groups']:>3d}g) "
              f"{s['test_images']:>5d}({s['test_groups']:>3d}g) "
              f"{s['train_malignant_pct']:>5.1f}/{s['val_malignant_pct']:>4.1f}/"
              f"{s['test_malignant_pct']:>4.1f}  {s['max_class_ratio_deviation_pp']:>4.1f}")
    print("-" * 100)
    tt = {k: sum(s[k] for s in cl) for k in
          ("train_images", "val_images", "test_images", "train_groups",
           "val_groups", "test_groups", "train_malignant", "val_malignant",
           "test_malignant", "total_images")}
    print(f"{'':4s} {'TOTAL':11s} {'':8s} "
          f"{tt['train_images']:>5d}({tt['train_groups']:>4d}g) "
          f"{tt['val_images']:>5d}({tt['val_groups']:>3d}g) "
          f"{tt['test_images']:>5d}({tt['test_groups']:>3d}g)")
    print(f"       {tt['train_images'] / tt['total_images'] * 100:.1f}% / "
          f"{tt['val_images'] / tt['total_images'] * 100:.1f}% / "
          f"{tt['test_images'] / tt['total_images'] * 100:.1f}%")
    small = [s for s in cl if s["test_images"] < 20]
    if small:
        print("\n  NOTE: a small test partition is a property of the client, not of "
              "the split:")
        for s in small:
            print(f"     {s['dataset']}: {s['test_images']} test image(s) from "
                  f"{s['test_groups']} group(s); the client holds "
                  f"{s['total_images']} image(s) in {s['total_groups']} group(s)")

    write_csv(out / "BMC_SPLITS_MANIFEST.csv", SPLIT_FIELDS, records)
    write_xlsx(out / "BMC_SPLITS_MANIFEST.xlsx", SPLIT_FIELDS, records,
               wide=("canonical_image_abspath", "canonical_mask_abspath",
                     "usage_policy", "group_provenance"))
    write_csv(out / "TABLE_4_split_summary.csv", SUMMARY_FIELDS, summary)
    write_xlsx(out / "TABLE_4_split_summary.xlsx", SUMMARY_FIELDS, summary,
               wide=("usage_policy",))
    (out / "BMC_SPLITS_REPORT.json").write_text(json.dumps({
        "phase": "R2_LEAKAGE_SAFE_PARTITIONING", "generated": now_iso(),
        "cohort_manifest": str(cohort_p), "seed": a.seed, "ratios": list(ratios),
        "assignment_rule": "GROUP_AWARE_CLASS_AWARE_DETERMINISTIC_V1",
        "unit": "the group key fixed by R1, assigned whole",
        "assignment_sha256": sha256_text(json.dumps(
            sorted((r["dataset"], r["group_key"], r["split"]) for r in records))),
        "clients": cl, "external": [s for s in summary if s["role"] == "EXTERNAL"],
        "totals": tt,
        "external_policy": EXTERNAL_POLICY,
        "group_key_leaks": leaks, "problems": problems,
        "checks": {d: ok for d, ok in checks},
        "canonical_files_touched": 0, "raw_touched": False,
        "verdict": "BMC_SPLITS_LOCKED" if not failed else "BLOCKED",
    }, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    print(f"WROTE {(out / 'BMC_SPLITS_REPORT.json').name}")
    print("\ncanonical files touched: NONE | cohort manifest: NONE | RAW: NONE")
    print("\nFINAL VERDICT:")
    if failed:
        print("BLOCKED")
        for d in failed:
            print(f"    FAILED: {d}")
        return 4
    print("BMC_SPLITS_LOCKED")
    print("The test partitions are now fixed and must not be consulted again until "
          "the final evaluation.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception:  # noqa: BLE001
        sys.stderr.write(traceback.format_exc())
        raise SystemExit(1)
