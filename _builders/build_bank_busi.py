#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_bank_busi.py
==================

PHASE 2B.3.1b - BUSI CANONICAL DATASET BANK BUILDER (patched, locked bindings)

Reference architecture for every later dataset builder.

LOCKED EVIDENCE BINDINGS
    pairing      BUSI\\audit\\03_pairing_audit.csv                      (647 rows)
                 image_filename, image_relative_path, class_folder,
                 number_of_matching_masks, matching_mask_filenames, pairing_status
    disposition  BUSI\\cluster_resolution\\04_busi_duplicate_group_map.csv
                 join on image_filename, disposition recommended_modeling_action,
                 group confirmed_duplicate_group_id, provenance future_split_group_key
    multi-mask   CANONICAL_DATASET_BANK\\_policies\\
                 BUSI_CANONICAL_MULTIMASK_POLICY_V1.json
                 required: policy_id == BUSI_CANONICAL_MULTIMASK_UNION_V1 and
                           policy_status == LOCKED_PHASE2B_CANONICAL_POLICY

TWO MODES
    --dry-run   validate everything; decode source pixels only to read native
                dimensions; write DRY_RUN-marked reports; write NO canonical PNG.
    --execute   permitted only when a PASSED dry-run state exists whose builder hash,
                schema hash and evidence hashes equal the current ones. Idempotent: an
                existing output file is left alone rather than rewritten.

NON-NEGOTIABLES
    RAW is opened for DECODING only and never written, renamed, moved or deleted. No
    audit is reopened, no duplicate analysis is rerun, no scientific policy is changed.
    No resize, crop, ROI, CLAHE, normalisation, denoising, sharpening, augmentation or
    model-specific processing. No split. No pseudo-mask: the union contains only pixels
    supported by source-provided masks.

FAIL-CLOSED
    Any count, join or policy mismatch stops the run and names the artifact and field
    that failed. Nothing is guessed.

Usage:
    python build_bank_busi.py --dry-run
    python build_bank_busi.py --execute
Requires: Python >= 3.8, Pillow, numpy. openpyxl optional (manifest XLSX).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import platform
import re
import sys
import traceback
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

DATASET = "BUSI"
DEFAULT_PREP_ROOT = (
    r"D:\PhD\002 FL Mul BC\01 ready\Springer\BMC Med Imaging"
    r"\Revision 1\dataset preparing"
)
DEFAULT_RAW_ROOT = r"H:\BC datasets\US"
BANK_NAME = "CANONICAL_DATASET_BANK"
ARCHIVE_DIR_NAMES = ("audit_old",)

# The BUSI source root is EVIDENCE-BOUND, not guessed. These are the only keys this
# builder will read it from, in order; the audit summary records the root the pairing
# table was generated against.
# set from --dataset-root; a filesystem location the operator supplies when the audit
# summary records none. It is still validated against every pairing row before use.
HUMAN_ROOT = {"path": ""}

ROOT_LOCATORS = ["source_path", "raw_root", "dataset_root", "source_root", "busi_root",
                 "raw_dataset_root", "source_dataset_root", "root"]

PAIRING_FILE = "03_pairing_audit.csv"
DISPOSITION_FILE = "04_busi_duplicate_group_map.csv"
SUMMARY_FILE = "16_busi_audit_summary.json"
CLUSTER_SUMMARY_FILE = "07_busi_cluster_resolution_summary.json"
POLICY_FILE = "BUSI_CANONICAL_MULTIMASK_POLICY_V1.json"
REQUIRED_POLICY_ID = "BUSI_CANONICAL_MULTIMASK_UNION_V1"
REQUIRED_POLICY_STATUS = "LOCKED_PHASE2B_CANONICAL_POLICY"

C_IMG_NAME = "image_filename"
C_IMG_PATH = "image_relative_path"
C_CLASS = "class_folder"
C_NMASK = "number_of_matching_masks"
C_MASKS = "matching_mask_filenames"
C_PAIRSTAT = "pairing_status"
C_DISP = "recommended_modeling_action"
C_DUPGRP = "confirmed_duplicate_group_id"
C_FUTKEY = "future_split_group_key"

MULTI_MASK_STATUS = "OK_MULTIPLE_MASKS"
D_UNIQUE = "ELIGIBLE_UNIQUE"
D_LOCK = "ELIGIBLE_GROUP_LOCK_REQUIRED"
D_EXCL = "EXCLUDE_CONFIRMED_LABEL_CONFLICT"
ELIGIBLE_DISPOSITIONS = (D_UNIQUE, D_LOCK)

# NOTE: no expected duplicate-group COUNT is asserted here. The scientifically correct
# invariant is derived from the authoritative table and the cluster summary at run time
# (see the duplicate-group semantics block), because "all confirmed duplicate groups"
# and "eligible group-lock groups" are different quantities.
EXPECT = {"source_images": 647, "source_masks": 665, "single_mask": 630,
          "multi_mask": 17, D_UNIQUE: 636, D_LOCK: 4, D_EXCL: 7, "eligible": 640}

CANONICAL_LABELS = {"BENIGN": 0, "MALIGNANT": 1}
SCHEMA_FIELDS = [
    "dataset", "canonical_record_id", "canonical_image_id", "canonical_image_path",
    "canonical_mask_id", "canonical_mask_path", "source_image_id", "source_mask_id",
    "source_image_path", "source_mask_path", "source_label", "canonical_binary_label",
    "patient_id", "case_id", "image_id", "lesion_id", "abnormality_id",
    "same_source_image_group", "duplicate_group_id", "classification_available",
    "segmentation_available", "image_conversion_status", "mask_conversion_status",
    "identity_provenance", "label_provenance", "mask_provenance",
    "source_spatial_height", "source_spatial_width", "canonical_spatial_height",
    "canonical_spatial_width", "source_image_format", "source_mask_format",
    "exclusion_or_special_status", "notes",
]
SPLIT_TOKENS = {"split", "splits", "fold", "folds", "kfold", "train", "training",
                "trainval", "val", "validation", "test", "testset", "testing",
                "holdout", "partition"}

SEP = " | "
LOGGER = logging.getLogger("bank_busi")


# ============================================================================
# Helpers
# ============================================================================


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def join_list(v: Sequence[Any], sep: str = SEP) -> str:
    return sep.join(str(x) for x in v)


def archived(p: Path) -> bool:
    return any(part.lower() in ARCHIVE_DIR_NAMES for part in p.parts[:-1])


def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as fh:
        for b in iter(lambda: fh.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def sha256_text(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def snapshot(root: Path) -> Dict[str, Tuple[int, int]]:
    """Filesystem metadata only; no file is opened."""
    out: Dict[str, Tuple[int, int]] = {}
    if not root.is_dir():
        return out
    for p in root.rglob("*"):
        try:
            if p.is_file():
                st = p.stat()
                out[str(p)] = (st.st_size, st.st_mtime_ns)
        except OSError:
            pass
    return out


def diff_state(a: dict, b: dict) -> dict:
    added = sorted(set(b) - set(a))
    removed = sorted(set(a) - set(b))
    modified = sorted(k for k in set(a) & set(b) if a[k] != b[k])
    return {"added": added, "removed": removed, "modified": modified,
            "unchanged": not (added or removed or modified)}


def is_split_like(path_str: str) -> bool:
    """Whole-token match with trailing digits stripped, so 'fold5' matches while
    'SAFETY_TESTS' and 'evaluation' do not. Attribution is by snapshot, not by name."""
    pth = Path(path_str)
    raw = [t for t in re.split(r"[^A-Za-z0-9]+", pth.stem) if t]
    raw += [t for part in pth.parts[:-1]
            for t in re.split(r"[^A-Za-z0-9]+", part) if t]
    toks = set()
    for t in raw:
        toks.add(t.lower())
        toks.add(re.sub(r"\d+$", "", t.lower()))
    return bool(toks & SPLIT_TOKENS)


def write_csv(path: Path, fields: Sequence[str], rows: Sequence[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(fields), extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)
    LOGGER.info("WROTE %-50s (%d data rows)", path.name, len(rows))


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    LOGGER.info("WROTE %-50s (%d chars)", path.name, len(text))


def read_rows(p: Path) -> List[dict]:
    with p.open(encoding="utf-8-sig", errors="replace") as fh:
        sample = fh.read(8192)
        fh.seek(0)
        delim = ";" if sample.count(";") > sample.count(",") else ","
        return list(csv.DictReader(fh, delimiter=delim))


def setup_logging(report_dir: Path, name: str) -> None:
    LOGGER.setLevel(logging.DEBUG)
    LOGGER.handlers.clear()
    f = logging.Formatter("%(asctime)s | %(levelname)-8s | %(message)s")
    sh = logging.StreamHandler(sys.stdout)
    sh.setFormatter(f)
    sh.setLevel(logging.INFO)
    LOGGER.addHandler(sh)
    report_dir.mkdir(parents=True, exist_ok=True)
    fh = logging.FileHandler(report_dir / name, mode="w", encoding="utf-8")
    fh.setFormatter(f)
    fh.setLevel(logging.DEBUG)
    LOGGER.addHandler(fh)


def find_one(files: Sequence[Path], filename: str) -> Tuple[Optional[Path], str]:
    hits = sorted(p for p in files if p.name.lower() == filename.lower())
    if len(hits) == 1:
        (only,) = hits
        return only, "UNIQUE"
    return None, ("MISSING" if not hits else f"AMBIGUOUS({len(hits)})")


def split_masks(cell: Any) -> List[str]:
    """matching_mask_filenames may hold several names. Never assume one."""
    if cell is None:
        return []
    return [x.strip() for x in re.split(r"[|;,]", str(cell)) if x.strip()]


def canonical_stem(class_folder: str, image_filename: str) -> str:
    """Deterministic, stable across reruns, order-independent, traceable, and with no
    invented identity. The hash is over the class-scoped source key."""
    key = f"{class_folder}/{image_filename}"
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(image_filename).stem).strip("_")
    return f"{safe or 'image'}__{sha256_text(key)[:10]}"


def binarize(arr, np) -> Tuple[Any, str]:
    """Canonical binary {0,255} only where the source supports it. A multilevel mask is
    refused rather than thresholded at a guessed cut."""
    vals = set(np.unique(arr).tolist())
    if vals <= {0}:
        return (arr * 0).astype(np.uint8), "ALL_ZERO_SOURCE"
    if vals <= {0, 255}:
        return arr.astype(np.uint8), "SOURCE_BINARY_0_255"
    if vals <= {0, 1}:
        return (arr * 255).astype(np.uint8), "SOURCE_BINARY_0_1_SCALED"
    return None, "MULTILEVEL_SOURCE_NOT_BINARISED"


# ============================================================================
# Evidence
# ============================================================================


def load_evidence(prep_root: Path) -> dict:
    ev: Dict[str, Any] = {"paths": {}, "hashes": {}, "problems": [], "status": {}}
    busi_dirs = [d for d in prep_root.iterdir()
                 if d.is_dir() and "busi" in d.name.lower().replace(" ", "")]
    files = [p for d in busi_dirs for p in d.rglob("*")
             if p.is_file() and not archived(p)]
    for key, fname, required in (("pairing", PAIRING_FILE, True),
                                 ("disposition", DISPOSITION_FILE, True),
                                 ("summary", SUMMARY_FILE, True),
                                 ("cluster_summary", CLUSTER_SUMMARY_FILE, False)):
        art, status = find_one(files, fname)
        ev["paths"][key] = art
        ev["status"][key] = status
        if art is not None:
            ev["hashes"][art.name] = sha256_file(art)
        elif required:
            ev["problems"].append(f"BLOCKED: {fname} -> {status}")

    # the source dataset root, read from the audit summary that the pairing table
    # was generated against
    ev["dataset_root"] = None
    ev["dataset_root_locator"] = ""
    if ev["paths"].get("summary"):
        try:
            blob = json.loads(ev["paths"]["summary"].read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001
            blob = {}
        for loc in ROOT_LOCATORS:
            val = blob.get(loc)
            if isinstance(val, str) and val.strip():
                ev["dataset_root"] = Path(val.strip())
                ev["dataset_root_locator"] = f"{SUMMARY_FILE} :: json:{loc}"
                break
        if ev["dataset_root"] is None and not HUMAN_ROOT["path"]:
            ev["problems"].append(
                f"BLOCKED: no source dataset root in {SUMMARY_FILE}; tried "
                + join_list([f"json:{k}" for k in ROOT_LOCATORS], ", ")
                + ". The root is never guessed. Supply it explicitly with "
                "--dataset-root \"H:\\BC datasets\\US\\<BUSI folder>\"; it is "
                "then verified against all 647 pairing rows before use.")
        if ev["dataset_root"] is None and HUMAN_ROOT["path"]:
            ev["dataset_root"] = Path(HUMAN_ROOT["path"])
            ev["dataset_root_locator"] = (
                "HUMAN_SUPPLIED --dataset-root, VERIFIED_AGAINST_ALL_PAIRING_ROWS "
                "(a filesystem location, not a scientific decision)")

    pol_dir = prep_root / BANK_NAME / "_policies"
    pol = pol_dir / POLICY_FILE
    ev["paths"]["policy"] = pol if pol.is_file() else None
    ev["status"]["policy"] = "UNIQUE" if pol.is_file() else "MISSING"
    ev["policy"] = None
    if pol.is_file():
        ev["hashes"][pol.name] = sha256_file(pol)
        try:
            blob = json.loads(pol.read_text(encoding="utf-8"))
        except Exception as exc:  # noqa: BLE001
            ev["problems"].append(f"BLOCKED: {POLICY_FILE} unreadable "
                                  f"({type(exc).__name__})")
            blob = {}
        pid, pst = blob.get("policy_id"), blob.get("policy_status")
        if pid != REQUIRED_POLICY_ID:
            ev["problems"].append(
                f"BLOCKED: policy_id is {pid!r}, required {REQUIRED_POLICY_ID!r}")
        elif pst != REQUIRED_POLICY_STATUS:
            ev["problems"].append(
                f"BLOCKED: policy_status is {pst!r}, required "
                f"{REQUIRED_POLICY_STATUS!r}")
        else:
            ev["policy"] = blob
    else:
        ev["problems"].append(
            f"BLOCKED: locked policy {POLICY_FILE} not found under {pol_dir}")
    return ev


# ============================================================================
# MAIN
# ============================================================================


def run(mode: str, prep_root: Path, raw_root: Path) -> int:
    bank_ds = prep_root / BANK_NAME / DATASET
    reports, manifests = bank_ds / "reports", bank_ds / "manifests"
    setup_logging(reports, f"busi_{mode}.log")
    LOGGER.info("=" * 78)
    LOGGER.info("PHASE 2B.3.1 BUSI CANONICAL BANK BUILDER  [%s]  (patched bindings)",
                mode.upper())
    LOGGER.info("=" * 78)
    LOGGER.info("script : %s", Path(__file__).resolve())
    LOGGER.info("python : %s", platform.python_version())
    LOGGER.info("RAW    : %s  [READ-ONLY: decoded, never written]", raw_root)
    LOGGER.info("bank   : %s", bank_ds)

    try:
        from PIL import Image
        import numpy as np
    except ImportError:
        LOGGER.critical("Pillow and numpy required: pip install pillow numpy")
        return 3
    Image.MAX_IMAGE_PIXELS = None

    before_raw = snapshot(raw_root)
    prep_dirs = [d for d in prep_root.iterdir()
                 if d.is_dir() and d.name != BANK_NAME]
    before_audit: Dict[str, Tuple[int, int]] = {}
    for d in prep_dirs:
        before_audit.update(snapshot(d))
    before_bank = snapshot(prep_root / BANK_NAME)

    # ---- 1. evidence ---------------------------------------------------------
    LOGGER.info("[1/8] locked evidence bindings ...")
    ev = load_evidence(prep_root)
    for k, st in ev["status"].items():
        art = ev["paths"].get(k)
        LOGGER.info("      %-13s %-12s %s", k, st, art.name if art else "")
    if ev["problems"]:
        for p in ev["problems"]:
            LOGGER.error("      %s", p)
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4
    pol = ev["policy"]
    LOGGER.info("      multi-mask policy = %s  [%s]", pol["policy_id"],
                pol["policy_status"])

    # ---- BUSI SOURCE-PATH DIAGNOSIS -----------------------------------------
    ds_root = ev["dataset_root"]
    LOGGER.info("")
    LOGGER.info("BUSI SOURCE-PATH DIAGNOSIS")
    LOGGER.info("  US parent root:    %s", raw_root)
    LOGGER.info("  BUSI dataset root: %s", ds_root)
    LOGGER.info("  root evidence:     %s", ev["dataset_root_locator"])

    pairing_rows = read_rows(ev["paths"]["pairing"])
    disp_rows = read_rows(ev["paths"]["disposition"])
    missing_cols = [c for c in (C_IMG_NAME, C_IMG_PATH, C_CLASS, C_NMASK, C_MASKS,
                                C_PAIRSTAT)
                    if pairing_rows and c not in pairing_rows[0]]
    missing_cols += [c for c in (C_IMG_NAME, C_DISP)
                     if disp_rows and c not in disp_rows[0]]
    if missing_cols:
        LOGGER.error("      declared column(s) absent: %s", join_list(missing_cols))
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4
    LOGGER.info("      pairing rows=%d  disposition rows=%d", len(pairing_rows),
                len(disp_rows))

    # Candidate bases come ONLY from evidence: the bound dataset root, and the CLI
    # parent as a declared fallback. Exactly one must resolve EVERY pairing row; this
    # is a deterministic base selection, never a recursive or fuzzy filename search.
    candidates = [("evidence-bound dataset root", ds_root)]
    if raw_root and raw_root != ds_root:
        candidates.append(("--raw-root parent", raw_root))
    resolving = []
    for label, base in candidates:
        if base is None:
            continue
        hits = sum(1 for r in pairing_rows
                   if (base / str(r[C_IMG_PATH]).strip()).is_file())
        LOGGER.info("  base candidate [%s] %s -> %d/%d image(s) resolve",
                    label, base, hits, len(pairing_rows))
        if hits == len(pairing_rows):
            resolving.append((label, base))
    if len(resolving) != 1:
        ex = pairing_rows[0] if pairing_rows else {}
        LOGGER.error("  example relative image: %s", ex.get(C_IMG_PATH, ""))
        for label, base in candidates:
            if base is not None:
                LOGGER.error("  example resolved [%s]: %s  exists=%s", label,
                             base / str(ex.get(C_IMG_PATH, "")),
                             (base / str(ex.get(C_IMG_PATH, ""))).is_file())
        LOGGER.error("  %s source base resolves every pairing row; the base is NOT "
                     "guessed and no fuzzy search is attempted",
                     "no" if not resolving else f"{len(resolving)} ambiguous")
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4
    base_label, source_base = resolving[0]
    ex = pairing_rows[0]
    ex_img = source_base / str(ex[C_IMG_PATH]).strip()
    ex_masks = [ex_img.parent / m for m in split_masks(ex[C_MASKS])]
    LOGGER.info("  selected base:     %s  [%s]", source_base, base_label)
    LOGGER.info("  example relative image: %s", ex[C_IMG_PATH])
    LOGGER.info("  example resolved image: %s", ex_img)
    LOGGER.info("  exists:            %s", ex_img.is_file())
    LOGGER.info("  example mask(s):   %s", join_list([str(m) for m in ex_masks], ", "))
    LOGGER.info("  exists:            %s",
                join_list([m.is_file() for m in ex_masks], ", "))

    # ---- 2. mask accounting ----------------------------------------------------
    LOGGER.info("[2/8] mask accounting (from %s) ...", PAIRING_FILE)
    n_single = sum(1 for r in pairing_rows if str(r[C_NMASK]).strip() == "1")
    n_multi = sum(1 for r in pairing_rows
                  if str(r[C_PAIRSTAT]).strip() == MULTI_MASK_STATUS)
    total_masks = 0
    for r in pairing_rows:
        try:
            total_masks += int(str(r[C_NMASK]).strip())
        except ValueError:
            pass
    LOGGER.info("      single-mask [%s == 1]            = %d", C_NMASK, n_single)
    LOGGER.info("      multi-mask  [%s == %s] = %d", C_PAIRSTAT, MULTI_MASK_STATUS,
                n_multi)
    LOGGER.info("      sum(%s)                          = %d", C_NMASK, total_masks)
    cross: Dict[str, Any] = {}
    if ev["paths"]["summary"]:
        blob = json.loads(ev["paths"]["summary"].read_text(encoding="utf-8"))
        cc = blob.get("class_counts") or {}
        cross = {"single_mask_pairs": blob.get("single_mask_pairs"),
                 "multiple_mask_pairs": blob.get("multiple_mask_pairs"),
                 "total_source_images": blob.get("total_source_images"),
                 "total_masks": blob.get("total_masks"),
                 "benign": cc.get("benign"), "malignant": cc.get("malignant")}
        LOGGER.info("      cross-check %s: %s", SUMMARY_FILE, cross)

    # ---- 3. crosswalk -----------------------------------------------------------
    LOGGER.info("[3/8] record-level crosswalk ...")
    p_names = [str(r[C_IMG_NAME]).strip() for r in pairing_rows]
    d_names = [str(r[C_IMG_NAME]).strip() for r in disp_rows]
    p_dupe = [k for k, v in Counter(p_names).items() if v > 1]
    d_dupe = [k for k, v in Counter(d_names).items() if v > 1]
    use_composite = False
    if p_dupe or d_dupe:
        if disp_rows and C_CLASS in disp_rows[0]:
            use_composite = True
            LOGGER.warning("      %s is not unique (%d pairing, %d disposition); using "
                           "the class-aware key (%s, %s), which BOTH artifacts support",
                           C_IMG_NAME, len(p_dupe), len(d_dupe), C_CLASS, C_IMG_NAME)
        else:
            LOGGER.error("      %s is not unique and %s is absent from %s; the join is "
                         "ambiguous and is NOT resolved silently",
                         C_IMG_NAME, C_CLASS, DISPOSITION_FILE)
            LOGGER.error("FINAL VERDICT:")
            LOGGER.error("BLOCKED")
            return 4

    def key_of(row: dict) -> str:
        if use_composite:
            return f"{str(row.get(C_CLASS, '')).strip()}/{str(row[C_IMG_NAME]).strip()}"
        return str(row[C_IMG_NAME]).strip()

    # future_split_group_key is the authoritative grouping key of the disposition
    # table. A key held by MORE THAN ONE record is a genuine leakage group; a key held
    # by exactly one record is a per-image singleton and is not a group. This is
    # derived by counting, never by matching on the key's name.
    futkey_counts = Counter(str(r.get(C_FUTKEY, "")).strip() for r in disp_rows
                            if str(r.get(C_FUTKEY, "")).strip())
    shared_futkeys = {k for k, n in futkey_counts.items() if n > 1}
    LOGGER.info("      %s: %d distinct value(s), %d shared by >1 record",
                C_FUTKEY, len(futkey_counts), len(shared_futkeys))

    disp_by_key = {key_of(r): r for r in disp_rows}
    joined: List[Tuple[dict, dict]] = []
    missing_disp: List[str] = []
    for r in pairing_rows:
        d = disp_by_key.get(key_of(r))
        if d is None:
            missing_disp.append(key_of(r))
        else:
            joined.append((r, d))
    LOGGER.info("      join key = %s",
                "(class_folder, image_filename)" if use_composite else "image_filename")
    LOGGER.info("      joins=%d  missing disposition=%d  ambiguous keys=%d",
                len(joined), len(missing_disp),
                0 if use_composite else len(p_dupe) + len(d_dupe))
    by_disp = Counter(str(d[C_DISP]).strip() for _p, d in joined)
    for k in (D_UNIQUE, D_LOCK, D_EXCL):
        LOGGER.info("         %-36s %d", k, by_disp.get(k, 0))

    # ---- 4. plan -----------------------------------------------------------------
    LOGGER.info("[4/8] planning canonical records ...")
    records: List[dict] = []
    for pr, dr in joined:
        disp = str(dr[C_DISP]).strip()
        is_elig = disp in ELIGIBLE_DISPOSITIONS
        cls = str(pr[C_CLASS]).strip()
        img_name = str(pr[C_IMG_NAME]).strip()
        img_rel = str(pr[C_IMG_PATH]).strip()
        masks = split_masks(pr[C_MASKS])
        label = ("BENIGN" if cls.lower().startswith("benign") else
                 "MALIGNANT" if cls.lower().startswith("malignant") else "UNRESOLVED")
        stem = canonical_stem(cls, img_name)
        folder = ("Benign" if label == "BENIGN" else
                  "Malignant" if label == "MALIGNANT" else "")
        img_out = f"{folder}/Images/{stem}.png" if is_elig and folder else ""
        msk_out = (f"{folder}/Masks/{stem}_mask.png"
                   if is_elig and folder and masks else "")
        # the pairing artifact supplies mask FILENAMES; the relative path is derived
        # from the image's own source folder and then validated by existence
        mask_rels = [str(Path(img_rel).parent / m) for m in masks]
        prov = ("SOURCE_NATIVE_SINGLE_MASK" if len(masks) == 1 else
                f"SOURCE_NATIVE_UNION_BY[{REQUIRED_POLICY_ID}]" if len(masks) > 1
                else "NO_VALID_GT")
        fut = str(dr.get(C_FUTKEY, "")).strip()
        confirmed = str(dr.get(C_DUPGRP, "")).strip()
        # the leakage group is the SHARED grouping key. confirmed_duplicate_group_id is
        # narrower - it marks confirmed duplicates only - so using it alone would leave
        # a provisionally locked pair with no lock at all.
        grp = fut if fut in shared_futkeys else ""
        records.append({
            "dataset": DATASET,
            "canonical_record_id": f"{DATASET}__{stem}",
            "canonical_image_id": f"{DATASET}__{stem}",
            "canonical_image_path": img_out,
            "canonical_mask_id": f"{DATASET}__{stem}_mask" if msk_out else "",
            "canonical_mask_path": msk_out,
            "source_image_id": img_name,
            "source_mask_id": join_list(masks, ","),
            "source_image_path": img_rel,
            "source_mask_path": join_list(mask_rels, ","),
            "source_label": label,
            "canonical_binary_label": (CANONICAL_LABELS.get(label, "UNRESOLVED")
                                       if is_elig else "NOT_IN_BINARY_TARGET"),
            "patient_id": "",
            "case_id": "",
            "image_id": img_name,
            "lesion_id": "",
            "abnormality_id": "",
            "same_source_image_group": f"{DATASET}__{stem}",
            "duplicate_group_id": grp,
            "classification_available": "YES" if is_elig else "NO",
            "segmentation_available": "YES" if is_elig and masks else "NO",
            "image_conversion_status": "PLANNED" if is_elig else "NOT_APPLICABLE",
            "mask_conversion_status": ("PLANNED" if msk_out else
                                       "NO_VALID_GT" if is_elig else "NOT_APPLICABLE"),
            "identity_provenance": "PATIENT_IDENTITY_UNAVAILABLE_IN_PUBLIC_RELEASE",
            "label_provenance": f"{PAIRING_FILE} :: csv:[{C_CLASS}]",
            "mask_provenance": prov,
            "source_spatial_height": "", "source_spatial_width": "",
            "canonical_spatial_height": "", "canonical_spatial_width": "",
            "source_image_format": Path(img_name).suffix.lstrip(".").upper(),
            "source_mask_format": (Path(masks[0]).suffix.lstrip(".").upper()
                                   if masks else ""),
            "exclusion_or_special_status": (
                "DUPLICATE_LOCKED" if disp == D_LOCK else
                "IN_BINARY_TARGET" if disp == D_UNIQUE else
                "LABEL_CONFLICT" if disp == D_EXCL else disp),
            "notes": (f"disposition={disp}; {C_NMASK}={len(masks)}; "
                      f"source_masks={join_list(masks, ',')}"
                      + (f"; {C_FUTKEY}={fut}" if fut else "")
                      + (f"; {C_DUPGRP}={confirmed}" if confirmed else "")
                      + (f"; {C_DUPGRP}=NONE_PROVISIONAL_LOCK_ONLY"
                         if grp and not confirmed else "")),
            "_disp": disp, "_eligible": is_elig, "_masks": mask_rels,
            "_nmask": len(masks),
        })
    eligible = [r for r in records if r["_eligible"]]
    excluded = [r for r in records if r["_disp"] == D_EXCL]
    multi_plans = [r for r in eligible if r["_nmask"] > 1]

    # ---- BUSI DUPLICATE-GROUP SEMANTICS --------------------------------------
    # "all confirmed duplicate groups" and "eligible group-lock groups" are DIFFERENT
    # quantities. Both are computed from the authoritative table; no group assignment
    # is changed, and no count is assumed.
    dup_records = [r for r in records if r["duplicate_group_id"]]
    dup_groups = sorted({r["duplicate_group_id"] for r in dup_records})
    lock_records = [r for r in records if r["_disp"] == D_LOCK]
    lock_with_grp = [r for r in lock_records if r["duplicate_group_id"]]
    lock_groups = sorted({r["duplicate_group_id"] for r in lock_with_grp})
    confirmed_records = [r for r in records if f"{C_DUPGRP}=" in r["notes"]
                         and "NONE_PROVISIONAL" not in r["notes"]]
    excl_with_grp = [r for r in excluded if r["duplicate_group_id"]]
    excl_groups = sorted({r["duplicate_group_id"] for r in excl_with_grp})
    uniq_with_grp = [r for r in records
                     if r["_disp"] == D_UNIQUE and r["duplicate_group_id"]]
    uniq_groups = sorted({r["duplicate_group_id"] for r in uniq_with_grp})
    # the invariant is EVIDENCE-BOUND: the cluster summary's group_lock_groups
    summary_lock_groups = None
    if ev["paths"].get("cluster_summary"):
        try:
            cs = json.loads(ev["paths"]["cluster_summary"].read_text(encoding="utf-8"))
            summary_lock_groups = cs.get("group_lock_groups")
        except Exception:  # noqa: BLE001
            summary_lock_groups = None
    LOGGER.info("")
    LOGGER.info("BUSI DUPLICATE-GROUP SEMANTICS")
    LOGGER.info("  records with duplicate_group_id (leakage groups): %d",
                len(dup_records))
    LOGGER.info("  records with a confirmed_duplicate_group_id:      %d",
                len(confirmed_records))
    LOGGER.info("  all distinct duplicate groups:          %d %s", len(dup_groups),
                dup_groups[:8])
    LOGGER.info("  eligible group-lock records:            %d", len(lock_records))
    LOGGER.info("  eligible group-lock groups:             %d %s", len(lock_groups),
                lock_groups[:8])
    LOGGER.info("  excluded conflict records:              %d", len(excluded))
    LOGGER.info("  excluded conflict groups:               %d %s", len(excl_groups),
                excl_groups[:8])
    LOGGER.info("  eligible-unique records carrying a group id: %d %s",
                len(uniq_with_grp), uniq_groups[:8])
    LOGGER.info("  cluster summary group_lock_groups:      %s", summary_lock_groups)
    if summary_lock_groups is not None and len(lock_groups) == summary_lock_groups:
        LOGGER.info("  explanation: group_lock_groups counts ONLY the groups spanned "
                    "by the %d %s records (=%d). The duplicate-group map additionally "
                    "carries %d group(s) belonging to other dispositions, so the table "
                    "shows %d records across %d groups. The two figures measure "
                    "different sets and do not conflict.",
                    len(lock_records), D_LOCK, len(lock_groups),
                    len(dup_groups) - len(lock_groups), len(dup_records),
                    len(dup_groups))
    else:
        LOGGER.warning("  explanation: the cluster summary reports %s group-lock "
                       "group(s) while the table yields %d among %s records. This is "
                       "NOT resolved by assumption; inspect both artifacts.",
                       summary_lock_groups, len(lock_groups), D_LOCK)
    LOGGER.info("")
    LOGGER.info("      eligible=%d  excluded retained=%d  union plans=%d",
                len(eligible), len(excluded), len(multi_plans))

    # ---- 5. resolve source files ---------------------------------------------------
    LOGGER.info("[5/8] resolving source files (decode for dimensions only) ...")
    img_ok = mask_ok = 0
    unresolved: List[str] = []
    for i, r in enumerate(records, 1):
        if not r["_eligible"]:
            continue
        sp = source_base / r["source_image_path"]
        if not sp.is_file():
            unresolved.append(f"image missing: {r['source_image_path']}")
            r["image_conversion_status"] = "UNRESOLVED_SOURCE"
            continue
        try:
            with Image.open(sp) as im:
                r["source_spatial_width"] = r["canonical_spatial_width"] = im.width
                r["source_spatial_height"] = r["canonical_spatial_height"] = im.height
            img_ok += 1
        except Exception as exc:  # noqa: BLE001
            unresolved.append(f"image undecodable: {r['source_image_path']} "
                              f"({type(exc).__name__})")
            r["image_conversion_status"] = "UNDECODABLE_SOURCE"
            continue
        gone = [m for m in r["_masks"] if not (source_base / m).is_file()]
        if r["_masks"] and not gone:
            mask_ok += 1
        elif gone:
            unresolved.append(f"mask missing: {join_list(gone, ',')}")
            r["mask_conversion_status"] = "UNRESOLVED_SOURCE"
        if i % 200 == 0:
            LOGGER.info("      %d/%d", i, len(records))
    LOGGER.info("      image mappings=%d  mask mappings=%d  unresolved=%d",
                img_ok, mask_ok, len(unresolved))
    for u in unresolved[:10]:
        LOGGER.error("         %s", u)

    # ---- 6. execute --------------------------------------------------------------
    png_written = 0
    if mode == "execute":
        LOGGER.info("[6/8] executing canonical conversion ...")
        for r in eligible:
            if r["image_conversion_status"] != "PLANNED":
                continue
            dst = bank_ds / r["canonical_image_path"]
            dst.parent.mkdir(parents=True, exist_ok=True)
            if not dst.exists():
                with Image.open(source_base / r["source_image_path"]) as im:
                    im.save(dst, format="PNG")      # no geometry change
                png_written += 1
            r["image_conversion_status"] = "CONVERTED"
            if not r["canonical_mask_path"]:
                continue
            acc = None
            enc = ""
            for m in r["_masks"]:
                with Image.open(source_base / m) as mi:
                    arr = np.array(mi.convert("L"))
                b, enc = binarize(arr, np)
                if b is None:
                    r["mask_conversion_status"] = "FAILED_MULTILEVEL_SOURCE"
                    acc = None
                    break
                acc = b if acc is None else np.maximum(acc, b)   # the locked union
            if acc is None:
                continue
            mdst = bank_ds / r["canonical_mask_path"]
            mdst.parent.mkdir(parents=True, exist_ok=True)
            if not mdst.exists():
                Image.fromarray(acc.astype(np.uint8), mode="L").save(mdst, format="PNG")
                png_written += 1
            r["mask_conversion_status"] = "CONVERTED"
            r["mask_provenance"] += f"; source_encoding={enc}"
    else:
        LOGGER.info("[6/8] dry-run: no canonical PNG is written")

    # ---- 7. tests ------------------------------------------------------------------
    LOGGER.info("[7/8] validation ...")
    after_bank = snapshot(prep_root / BANK_NAME)
    pre_splits = [k for k in before_bank if is_split_like(k)]
    new_splits = [k for k in after_bank if k not in before_bank and is_split_like(k)]
    mod_splits = [k for k in after_bank if k in before_bank
                  and after_bank[k] != before_bank[k] and is_split_like(k)]
    new_png = [k for k in after_bank
               if k not in before_bank and k.lower().endswith(".png")]
    raw_diff = diff_state(before_raw, snapshot(raw_root))
    after_audit: Dict[str, Tuple[int, int]] = {}
    for d in prep_dirs:
        after_audit.update(snapshot(d))
    aud_diff = diff_state(before_audit, after_audit)
    paths = [r["canonical_image_path"] for r in eligible if r["canonical_image_path"]]
    collisions = [n for n, c in Counter(paths).items() if c > 1]

    T: List[Tuple[str, str, Optional[bool], str]] = []

    def add(i: str, d: str, ok: Optional[bool], b: str) -> None:
        T.append((i, d, ok, b))

    add("T01", f"pairing rows == {EXPECT['source_images']}",
        len(pairing_rows) == EXPECT["source_images"], f"{len(pairing_rows)}")
    add("T02", f"disposition rows == {EXPECT['source_images']}",
        len(disp_rows) == EXPECT["source_images"], f"{len(disp_rows)}")
    add("T03", f"successful joins == {EXPECT['source_images']}",
        len(joined) == EXPECT["source_images"] and not missing_disp,
        f"{len(joined)} joined, {len(missing_disp)} missing")
    add("T04", "no ambiguous join key", (not (p_dupe or d_dupe)) or use_composite,
        f"pairing dupes={len(p_dupe)}, disposition dupes={len(d_dupe)}, "
        f"composite={use_composite}")
    add("T05", f"single-mask == {EXPECT['single_mask']}",
        n_single == EXPECT["single_mask"],
        f"{n_single} @ {PAIRING_FILE}::[{C_NMASK}=1]")
    add("T06", f"multi-mask == {EXPECT['multi_mask']}",
        n_multi == EXPECT["multi_mask"],
        f"{n_multi} @ {PAIRING_FILE}::[{C_PAIRSTAT}={MULTI_MASK_STATUS}]")
    add("T07", f"source masks == {EXPECT['source_masks']}",
        total_masks == EXPECT["source_masks"], f"sum({C_NMASK})={total_masks}")
    add("T08", f"{D_UNIQUE} == {EXPECT[D_UNIQUE]}",
        by_disp.get(D_UNIQUE) == EXPECT[D_UNIQUE], f"{by_disp.get(D_UNIQUE)}")
    add("T09", f"{D_LOCK} == {EXPECT[D_LOCK]}",
        by_disp.get(D_LOCK) == EXPECT[D_LOCK], f"{by_disp.get(D_LOCK)}")
    add("T10", f"{D_EXCL} == {EXPECT[D_EXCL]}",
        by_disp.get(D_EXCL) == EXPECT[D_EXCL], f"{by_disp.get(D_EXCL)}")
    add("T11", f"eligible == {EXPECT['eligible']}",
        len(eligible) == EXPECT["eligible"], f"{len(eligible)}")
    add("T12", "excluded records retained without canonical paths",
        len(excluded) == EXPECT[D_EXCL]
        and all(not r["canonical_image_path"] for r in excluded),
        f"{len(excluded)} retained, "
        f"{sum(1 for r in excluded if r['canonical_image_path'])} with a path")
    add("T13", "every eligible image resolved", img_ok == len(eligible),
        f"{img_ok}/{len(eligible)}")
    add("T14", "every eligible segmentation target resolved", mask_ok == len(eligible),
        f"{mask_ok}/{len(eligible)}")
    add("T15", f"multi-mask union plans == {EXPECT['multi_mask']}",
        len(multi_plans) == EXPECT["multi_mask"], f"{len(multi_plans)}")
    add("T16", "locked policy id and status verified",
        pol["policy_id"] == REQUIRED_POLICY_ID
        and pol["policy_status"] == REQUIRED_POLICY_STATUS,
        f"{pol['policy_id']} / {pol['policy_status']}")
    add("T17", "no patient ID invented", all(not r["patient_id"] for r in records),
        "patient_id empty on every record")
    add("T18", f"group-lock records == {EXPECT[D_LOCK]}",
        len([r for r in records if r["_disp"] == D_LOCK]) == EXPECT[D_LOCK],
        f"{len([r for r in records if r['_disp'] == D_LOCK])}")
    add("T19", "eligible group-lock groups match the cluster summary",
        (summary_lock_groups is None) or (len(lock_groups) == summary_lock_groups),
        f"table groups among {D_LOCK} = {len(lock_groups)} {lock_groups}; "
        f"{CLUSTER_SUMMARY_FILE}::group_lock_groups = {summary_lock_groups}")
    add("T19b", f"every {D_LOCK} record carries a duplicate group id",
        len(lock_with_grp) == len(lock_records),
        f"{len(lock_with_grp)}/{len(lock_records)}")
    add("T19c", "all-confirmed-duplicate-groups reported separately from group-lock "
        "groups", True,
        f"all groups={len(dup_groups)} over {len(dup_records)} record(s); "
        f"group-lock groups={len(lock_groups)}; excluded-conflict groups="
        f"{len(excl_groups)}; eligible-unique carrying a group={len(uniq_with_grp)}")
    add("T20", "duplicate group preserved and not used as identity",
        bool(dup_records) and all(not r["patient_id"] for r in dup_records),
        f"{len(dup_records)} record(s) carry a group id, none carries a patient id")
    add("T21", "no filename collision", not collisions, f"{len(collisions)}")
    add("T22", "no unresolved eligible record", not unresolved, f"{len(unresolved)}")
    add("T23", "RAW unchanged", raw_diff["unchanged"],
        f"modified {len(raw_diff['modified'])}, added {len(raw_diff['added'])}")
    add("T24", "completed audit artifacts unchanged", aud_diff["unchanged"],
        f"modified {len(aud_diff['modified'])}")
    add("T25", "no split created or modified by this run",
        not new_splits and not mod_splits,
        f"preexisting={len(pre_splits)}, new={len(new_splits)}, "
        f"modified={len(mod_splits)}")
    add("T26", "dry-run wrote no canonical PNG",
        (mode != "dry-run") or not new_png, f"{len(new_png)} new PNG")
    add("T27", "no split field in the manifest schema",
        not any(is_split_like(f) for f in SCHEMA_FIELDS), "schema checked")
    add("T28", "no preprocessing in the conversion path", None,
        "code-scope declaration: decode, optional union, lossless PNG; no resize, "
        "crop, ROI, CLAHE, normalisation or augmentation call exists")
    verified = [t for t in T if t[2] is not None]
    failed = [t for t in verified if not t[2]]
    for t in T:
        (LOGGER.info if t[2] is not False else LOGGER.error)(
            "      %-5s [%-11s] %-50s %s", t[0],
            "DECLARATION" if t[2] is None else ("PASS" if t[2] else "FAIL"),
            t[1][:50], t[3][:50])

    # ---- 8. outputs -------------------------------------------------------------------
    LOGGER.info("[8/8] writing outputs ...")
    tag = "DRY_RUN_" if mode == "dry-run" else ""
    out_rows = [{k: r[k] for k in SCHEMA_FIELDS} for r in records]
    write_csv(manifests / f"{tag}canonical_manifest.csv", SCHEMA_FIELDS, out_rows)
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "canonical_manifest"
        ws.append(SCHEMA_FIELDS)
        for r in out_rows:
            ws.append([r.get(h, "") for h in SCHEMA_FIELDS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = (f"A1:{get_column_letter(len(SCHEMA_FIELDS))}"
                              f"{len(out_rows) + 1}")
        wb.save(manifests / f"{tag}canonical_manifest.xlsx")
        LOGGER.info("WROTE %-50s (1 sheet)", f"{tag}canonical_manifest.xlsx")
    except ImportError:
        LOGGER.warning("openpyxl not installed; manifest XLSX skipped")

    blocked = bool(failed)
    state = {
        "phase": "2B.3.1b", "dataset": DATASET, "mode": mode, "generated": now_iso(),
        "builder_sha256": sha256_file(Path(__file__).resolve()),
        "schema_sha256": sha256_text(join_list(SCHEMA_FIELDS, ",")),
        "evidence_hashes": ev["hashes"],
        "policy": {"policy_id": pol["policy_id"],
                   "policy_status": pol["policy_status"]},
        "counts": {"pairing_rows": len(pairing_rows),
                   "disposition_rows": len(disp_rows), "joins": len(joined),
                   "eligible": len(eligible), "excluded": len(excluded),
                   "single_mask": n_single, "multi_mask": n_multi,
                   "source_masks": total_masks, "union_plans": len(multi_plans),
                   "records_with_duplicate_group_id": len(dup_records),
                   "all_distinct_duplicate_groups": len(dup_groups),
                   "eligible_group_lock_records": len(lock_records),
                   "eligible_group_lock_groups": len(lock_groups),
                   "excluded_conflict_records": len(excluded),
                   "excluded_conflict_groups": len(excl_groups),
                   "eligible_unique_records_with_group": len(uniq_with_grp),
                   "cluster_summary_group_lock_groups": summary_lock_groups},
        "image_mappings_resolved": img_ok, "mask_mappings_resolved": mask_ok,
        "unresolved": unresolved[:50], "filename_collisions": collisions,
        "canonical_png_written": png_written,
        "preexisting_split_artifacts": len(pre_splits),
        "newly_created_split_artifacts": len(new_splits),
        "modified_split_artifacts": len(mod_splits),
        "raw_modified": not raw_diff["unchanged"],
        "audit_modified": not aud_diff["unchanged"],
        "tests_pass": sum(1 for t in verified if t[2]), "tests_fail": len(failed),
        "verdict": ("BLOCKED" if blocked else
                    "READY_FOR_BUSI_CANONICAL_EXECUTION" if mode == "dry-run"
                    else "BUSI_CANONICAL_BUILD_COMPLETE"),
    }
    write_text(reports / f"{tag}BUSI_BUILD_STATE.json",
               json.dumps(state, indent=2, ensure_ascii=False, default=str))

    LOGGER.info("=" * 70)
    LOGGER.info("PHASE 2B.3.1 BUSI CANONICAL BANK - %s",
                "PATCHED DRY RUN" if mode == "dry-run" else "EXECUTION")
    LOGGER.info("")
    LOGGER.info("Evidence bindings:")
    LOGGER.info("  Pairing:           %s (%d rows)", PAIRING_FILE, len(pairing_rows))
    LOGGER.info("  Disposition:       %s (%d rows)", DISPOSITION_FILE, len(disp_rows))
    LOGGER.info("  Multi-mask policy: %s [%s]", pol["policy_id"], pol["policy_status"])
    LOGGER.info("")
    LOGGER.info("Source accounting:")
    LOGGER.info("  %s / %s / %s / %s", len(pairing_rows), cross.get("benign", "?"),
                cross.get("malignant", "?"), total_masks)
    LOGGER.info("  %d single-mask", n_single)
    LOGGER.info("  %d multi-mask", n_multi)
    LOGGER.info("")
    LOGGER.info("Disposition:")
    LOGGER.info("  %d unique", by_disp.get(D_UNIQUE, 0))
    LOGGER.info("  %d group-lock", by_disp.get(D_LOCK, 0))
    LOGGER.info("  %d excluded", by_disp.get(D_EXCL, 0))
    LOGGER.info("")
    LOGGER.info("Crosswalk:")
    LOGGER.info("  pairing rows:      %d", len(pairing_rows))
    LOGGER.info("  disposition rows:  %d", len(disp_rows))
    LOGGER.info("  successful joins:  %d", len(joined))
    LOGGER.info("  ambiguous joins:   %d",
                0 if use_composite else len(p_dupe) + len(d_dupe))
    LOGGER.info("")
    LOGGER.info("Canonical plan:")
    LOGGER.info("  eligible records:  %d", len(eligible))
    LOGGER.info("  image mappings:    %d", img_ok)
    LOGGER.info("  mask mappings:     %d", mask_ok)
    LOGGER.info("  multi-mask union plans: %d", len(multi_plans))
    LOGGER.info("  excluded retained in manifest: %d", len(excluded))
    LOGGER.info("")
    LOGGER.info("Grouping:")
    LOGGER.info("  patient IDs:       0 (unavailable in the public release)")
    LOGGER.info("  group-lock records: %d",
                len([r for r in records if r["_disp"] == D_LOCK]))
    LOGGER.info("  all duplicate-group records: %d across %d group(s)",
                len(dup_records), len(dup_groups))
    LOGGER.info("  eligible group-lock groups:  %d %s", len(lock_groups), lock_groups)
    LOGGER.info("  excluded-conflict groups:    %d %s", len(excl_groups), excl_groups)
    LOGGER.info("")
    LOGGER.info("Safety:")
    LOGGER.info("  RAW modified:      %s", state["raw_modified"])
    LOGGER.info("  audit modified:    %s", state["audit_modified"])
    LOGGER.info("  canonical PNG written: %d", png_written)
    LOGGER.info("  splits created:    %d (pre-existing %d, modified %d)",
                len(new_splits), len(pre_splits), len(mod_splits))
    LOGGER.info("")
    LOGGER.info("TESTS PASS: %d", state["tests_pass"])
    LOGGER.info("TESTS FAIL: %d", state["tests_fail"])
    LOGGER.info("")
    LOGGER.info("FINAL VERDICT:")
    if blocked:
        LOGGER.error("BLOCKED")
        LOGGER.error("=" * 70)
        for t in failed:
            LOGGER.error("    FAILED %s: %s | %s", t[0], t[1], t[3])
        return 4
    if mode == "dry-run":
        LOGGER.info("READY_FOR_BUSI_CANONICAL_EXECUTION")
        LOGGER.info("=" * 70)
        LOGGER.info("STOP. --execute is not run automatically and needs human "
                    "approval.")
    else:
        LOGGER.info("BUSI_CANONICAL_BUILD_COMPLETE")
    return 0


def execute_gate(prep_root: Path) -> Tuple[bool, str]:
    """--execute requires a PASSED dry-run whose builder, schema and evidence hashes
    equal the current ones. A stale approval never authorises execution."""
    st = prep_root / BANK_NAME / DATASET / "reports" / "DRY_RUN_BUSI_BUILD_STATE.json"
    if not st.is_file():
        return False, f"no dry-run state at {st}"
    try:
        d = json.loads(st.read_text(encoding="utf-8"))
    except Exception as exc:  # noqa: BLE001
        return False, f"dry-run state unreadable ({type(exc).__name__})"
    if d.get("verdict") != "READY_FOR_BUSI_CANONICAL_EXECUTION":
        return False, f"the last dry-run verdict was {d.get('verdict')!r}"
    if d.get("builder_sha256") != sha256_file(Path(__file__).resolve()):
        return False, "the builder changed since the approved dry-run; rerun --dry-run"
    if d.get("schema_sha256") != sha256_text(join_list(SCHEMA_FIELDS, ",")):
        return False, "the schema changed since the approved dry-run; rerun --dry-run"
    ev = load_evidence(prep_root)
    for name, h in (d.get("evidence_hashes") or {}).items():
        if ev["hashes"].get(name) != h:
            return False, f"evidence artifact {name} changed since the dry-run"
    return True, "the dry-run approval matches the current builder, schema and evidence"


def main(argv: Optional[Sequence[str]] = None) -> int:
    p = argparse.ArgumentParser(description="BUSI canonical bank builder.")
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--dry-run", action="store_true")
    g.add_argument("--execute", action="store_true")
    p.add_argument("--prep-root", default=DEFAULT_PREP_ROOT)
    p.add_argument("--raw-root", default=DEFAULT_RAW_ROOT)
    p.add_argument("--dataset-root", default="",
                   help="the BUSI source folder, used only when the audit summary "
                        "records no root; validated against all pairing rows")
    a = p.parse_args(argv)
    HUMAN_ROOT["path"] = a.dataset_root
    mode = "dry-run" if a.dry_run else "execute"
    if mode == "execute":
        ok, why = execute_gate(Path(a.prep_root))
        if not ok:
            sys.stderr.write(f"EXECUTE refused: {why}\n")
            return 5
        sys.stderr.write(f"execute gate: {why}\n")
    try:
        return run(mode, Path(a.prep_root), Path(a.raw_root))
    except Exception:  # noqa: BLE001
        if LOGGER.handlers:
            LOGGER.critical("UNHANDLED EXCEPTION\n%s", traceback.format_exc())
        else:
            sys.stderr.write(traceback.format_exc())
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
