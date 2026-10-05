#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_bank_bus_uc.py
====================

PHASE 2B.3.2 - BUS-UC CANONICAL DATASET BANK BUILDER

Same fail-closed architecture as the BUSI builder. Nothing is carried over from BUSI's
science: BUS-UC has ONE mask per image, so no union is ever formed, and its threshold
comes from its own locked policy.

LOCKED EVIDENCE BINDINGS
    manifest    BUS_UC\\audit\\12_bus_uc_audit_manifest.csv          (811 rows)
                join key sample_key - class-scoped, because image_filename repeats
                across class folders
                image_path, mask_paths, class_label, mask_count, pairing_status,
                patient_id_status, near_duplicate_candidate_flag, exact_duplicate_group
    integrity   BUS_UC\\audit\\19_bus_uc_mask_integrity_cases.csv     (3 rows)
                the excluded sample_keys are READ FROM THIS ARTIFACT, never hard-coded
    root        16_bus_uc_audit_summary.json :: json:raw_root
    policy      CANONICAL_DATASET_BANK\\_policies\\
                BUS_UC_CANONICAL_MASK_POLICY_V1.json
                required policy_id BUS_UC_CANONICAL_MASK_BINARIZATION_V1 and
                policy_status LOCKED_PHASE2B_CANONICAL_POLICY; the threshold is READ
                FROM THE POLICY and is not a constant in this file

MASK RULE (from the policy, applied verbatim)
    verify R = G = B per pixel -> reduce to one grayscale plane ->
    canonical = 255 * (gray >= threshold) -> single-channel lossless PNG
    No resize, resample, crop, pad, geometric transform, pseudo-mask or substitution.

THE THREE INTEGRITY RECORDS are excluded from the canonical joint cohort, retained in
the manifest with source truth and an explicit reason, and written no canonical path.

Usage:
    python build_bank_bus_uc.py --dry-run
    python build_bank_bus_uc.py --execute
Requires: Python >= 3.8, Pillow, numpy. openpyxl optional.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import platform
import re
import sys
import traceback
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

DATASET = "BUS_UC"
DEFAULT_PREP_ROOT = (
    r"D:\PhD\002 FL Mul BC\01 ready\Springer\BMC Med Imaging"
    r"\Revision 1\dataset preparing"
)
BANK_NAME = "CANONICAL_DATASET_BANK"
ARCHIVE_DIR_NAMES = ("audit_old",)

MANIFEST_FILE = "12_bus_uc_audit_manifest.csv"
INTEGRITY_FILE = "19_bus_uc_mask_integrity_cases.csv"
SUMMARY_FILE = "16_bus_uc_audit_summary.json"
POLICY_FILE = "BUS_UC_CANONICAL_MASK_POLICY_V1.json"
REQUIRED_POLICY_ID = "BUS_UC_CANONICAL_MASK_BINARIZATION_V1"
REQUIRED_POLICY_STATUS = "LOCKED_PHASE2B_CANONICAL_POLICY"
ROOT_LOCATORS = ["raw_root", "source_path", "dataset_root", "source_root"]

C_KEY = "sample_key"
C_IMGNAME = "image_filename"
C_IMGPATH = "image_path"
C_MASKNAMES = "mask_filenames"
C_MASKPATHS = "mask_paths"
C_CLASS = "class_label"
C_CLASSSRC = "class_label_source"
C_NMASK = "mask_count"
C_PAIRSTAT = "pairing_status"
C_PIDSTAT = "patient_id_status"
C_CASE = "case_id"
C_LESION = "lesion_id"
C_NEARDUP = "near_duplicate_candidate_flag"
C_EXACTDUP = "exact_duplicate_group"
C_ENC = "mask_encoding"
C_W, C_H = "width", "height"

EXPECT = {"records": 811, "benign": 358, "malignant": 453,
          "eligible": 808, "eligible_benign": 357, "eligible_malignant": 451,
          "excluded": 3, "png": 1616}

# exclusion status is DERIVED from the integrity artifact's own flags
EXCLUSION_RULES = [
    ("EMPTY_ALL_ZERO_MASK", None, "EXCLUDE_EMPTY_SEGMENTATION_MASK"),
    ("IMAGE_MASK_DIMENSION_MISMATCH", "SYSTEMATIC_SCALE_DIFFERENCE",
     "EXCLUDE_MASK_SCALE_DIMENSION_MISMATCH"),
    ("IMAGE_MASK_DIMENSION_MISMATCH", None, "EXCLUDE_MASK_DIMENSION_MISMATCH"),
]

CANONICAL_LABELS = {"BENIGN": 0, "MALIGNANT": 1}
SCHEMA_FIELDS = [
    "dataset", "canonical_record_id", "sample_key", "source_image_id",
    "source_mask_id", "source_image_path", "source_mask_path", "source_label",
    "canonical_binary_label", "canonical_image_path", "canonical_mask_path",
    "patient_id", "case_id", "lesion_id",
    "classification_available", "segmentation_available",
    "image_conversion_status", "mask_conversion_status",
    "identity_provenance", "label_provenance", "mask_provenance",
    "source_image_height", "source_image_width",
    "source_mask_height", "source_mask_width",
    "canonical_image_height", "canonical_image_width",
    "canonical_mask_height", "canonical_mask_width",
    "source_mask_encoding", "canonical_mask_encoding",
    "canonical_binarization_threshold", "canonical_binarization_rule",
    "threshold_provenance", "exclusion_status", "exclusion_reason", "notes",
]
SPLIT_TOKENS = {"split", "splits", "fold", "folds", "kfold", "train", "training",
                "trainval", "val", "validation", "test", "testset", "testing",
                "holdout", "partition"}

SEP = " | "
LOGGER = logging.getLogger("bank_bus_uc")


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def join_list(v: Sequence[Any], sep: str = SEP) -> str:
    return sep.join(str(x) for x in v)


def archived(p: Path) -> bool:
    return any(x.lower() in ARCHIVE_DIR_NAMES for x in p.parts[:-1])


def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as fh:
        for b in iter(lambda: fh.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def sha256_text(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def snapshot(root: Path) -> Dict[str, Tuple[int, int]]:
    out: Dict[str, Tuple[int, int]] = {}
    if not root.is_dir():
        return out
    for p in root.rglob("*"):
        try:
            if p.is_file():
                st = p.stat()
                out[str(p)] = (st.st_size, st.st_mtime_ns)
        except OSError:
            pass
    return out


def diff_state(a: dict, b: dict) -> dict:
    added = sorted(set(b) - set(a))
    removed = sorted(set(a) - set(b))
    modified = sorted(k for k in set(a) & set(b) if a[k] != b[k])
    return {"added": added, "removed": removed, "modified": modified,
            "unchanged": not (added or removed or modified)}


def is_split_like(path_str: str) -> bool:
    pth = Path(path_str)
    raw = [t for t in re.split(r"[^A-Za-z0-9]+", pth.stem) if t]
    raw += [t for part in pth.parts[:-1]
            for t in re.split(r"[^A-Za-z0-9]+", part) if t]
    toks = set()
    for t in raw:
        toks.add(t.lower())
        toks.add(re.sub(r"\d+$", "", t.lower()))
    return bool(toks & SPLIT_TOKENS)


def write_csv(path: Path, fields: Sequence[str], rows: Sequence[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(fields), extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)
    LOGGER.info("WROTE %-48s (%d data rows)", path.name, len(rows))


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    LOGGER.info("WROTE %-48s (%d chars)", path.name, len(text))


def read_rows(p: Path) -> List[dict]:
    with p.open(encoding="utf-8-sig", errors="replace") as fh:
        s = fh.read(8192)
        fh.seek(0)
        d = ";" if s.count(";") > s.count(",") else ","
        return list(csv.DictReader(fh, delimiter=d))


def setup_logging(report_dir: Path, name: str) -> None:
    LOGGER.setLevel(logging.DEBUG)
    LOGGER.handlers.clear()
    f = logging.Formatter("%(asctime)s | %(levelname)-8s | %(message)s")
    sh = logging.StreamHandler(sys.stdout)
    sh.setFormatter(f)
    sh.setLevel(logging.INFO)
    LOGGER.addHandler(sh)
    report_dir.mkdir(parents=True, exist_ok=True)
    fh = logging.FileHandler(report_dir / name, mode="w", encoding="utf-8")
    fh.setFormatter(f)
    fh.setLevel(logging.DEBUG)
    LOGGER.addHandler(fh)


def find_one(files: Sequence[Path], filename: str) -> Tuple[Optional[Path], str]:
    hits = sorted(p for p in files if p.name.lower() == filename.lower())
    if len(hits) == 1:
        (only,) = hits
        return only, "UNIQUE"
    return None, ("MISSING" if not hits else f"AMBIGUOUS({len(hits)})")


def canonical_stem(sample_key: str) -> str:
    """Deterministic, stable, order-independent, traceable. sample_key is already
    class-scoped, so the hash is taken over it directly."""
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", sample_key.replace("::", "__")).strip("_")
    safe = re.sub(r"\.png$|\.PNG$", "", safe)
    return f"{safe or 'image'}__{sha256_text(sample_key)[:10]}"


def exclusion_for(flags: str, kind: str) -> str:
    for flag, want_kind, status in EXCLUSION_RULES:
        if flag in flags and (want_kind is None or want_kind == kind):
            return status
    return "EXCLUDE_MASK_INTEGRITY_UNCLASSIFIED"


def load_evidence(prep_root: Path) -> dict:
    ev: Dict[str, Any] = {"paths": {}, "status": {}, "hashes": {}, "problems": []}
    dirs = [d for d in prep_root.iterdir() if d.is_dir()
            and "busuc" in d.name.lower().replace(" ", "").replace("_", "")]
    files = [p for d in dirs for p in d.rglob("*") if p.is_file() and not archived(p)]
    for key, fname in (("manifest", MANIFEST_FILE), ("integrity", INTEGRITY_FILE),
                       ("summary", SUMMARY_FILE)):
        art, st = find_one(files, fname)
        ev["paths"][key] = art
        ev["status"][key] = st
        if art is not None:
            ev["hashes"][art.name] = sha256_file(art)
        else:
            ev["problems"].append(f"BLOCKED: {fname} -> {st}")

    ev["dataset_root"] = None
    ev["dataset_root_locator"] = ""
    if ev["paths"]["summary"]:
        blob = json.loads(ev["paths"]["summary"].read_text(encoding="utf-8"))
        for loc in ROOT_LOCATORS:
            v = blob.get(loc)
            if isinstance(v, str) and v.strip():
                ev["dataset_root"] = Path(v.strip())
                ev["dataset_root_locator"] = f"{SUMMARY_FILE} :: json:{loc}"
                break
        if ev["dataset_root"] is None:
            ev["problems"].append(
                f"BLOCKED: no source root in {SUMMARY_FILE}; tried "
                + join_list([f"json:{k}" for k in ROOT_LOCATORS], ", "))

    pol = prep_root / BANK_NAME / "_policies" / POLICY_FILE
    ev["paths"]["policy"] = pol if pol.is_file() else None
    ev["status"]["policy"] = "UNIQUE" if pol.is_file() else "MISSING"
    ev["policy"] = None
    if pol.is_file():
        ev["hashes"][pol.name] = sha256_file(pol)
        blob = json.loads(pol.read_text(encoding="utf-8"))
        if blob.get("policy_id") != REQUIRED_POLICY_ID:
            ev["problems"].append(f"BLOCKED: policy_id is {blob.get('policy_id')!r}")
        elif blob.get("policy_status") != REQUIRED_POLICY_STATUS:
            ev["problems"].append(
                f"BLOCKED: policy_status is {blob.get('policy_status')!r}")
        elif not isinstance(blob.get("canonical_threshold"), int):
            ev["problems"].append("BLOCKED: canonical_threshold missing or not an "
                                  "integer in the locked policy")
        else:
            ev["policy"] = blob
    else:
        ev["problems"].append(f"BLOCKED: locked policy {POLICY_FILE} not found. Run "
                              "lock_bus_uc_mask_policy.py first.")
    return ev


def run(mode: str, prep_root: Path) -> int:
    bank_ds = prep_root / BANK_NAME / DATASET
    reports, manifests = bank_ds / "reports", bank_ds / "manifests"
    setup_logging(reports, f"bus_uc_{mode}.log")
    LOGGER.info("=" * 78)
    LOGGER.info("PHASE 2B.3.2 BUS-UC CANONICAL BANK BUILDER  [%s]", mode.upper())
    LOGGER.info("=" * 78)
    LOGGER.info("script : %s", Path(__file__).resolve())
    LOGGER.info("python : %s", platform.python_version())
    LOGGER.info("bank   : %s", bank_ds)

    try:
        from PIL import Image
        import numpy as np
    except ImportError:
        LOGGER.critical("Pillow and numpy required: pip install pillow numpy")
        return 3
    Image.MAX_IMAGE_PIXELS = None

    prep_dirs = [d for d in prep_root.iterdir()
                 if d.is_dir() and d.name != BANK_NAME]
    before_audit: Dict[str, Tuple[int, int]] = {}
    for d in prep_dirs:
        before_audit.update(snapshot(d))
    before_bank = snapshot(prep_root / BANK_NAME)

    # ---- 1. evidence -------------------------------------------------------
    LOGGER.info("[1/8] locked evidence bindings ...")
    ev = load_evidence(prep_root)
    for k, st in ev["status"].items():
        art = ev["paths"].get(k)
        LOGGER.info("      %-10s %-12s %s", k, st, art.name if art else "")
    if ev["problems"]:
        for p in ev["problems"]:
            LOGGER.error("      %s", p)
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4
    pol = ev["policy"]
    THRESHOLD = int(pol["canonical_threshold"])
    RULE = pol.get("canonical_rule", "")
    src_root = ev["dataset_root"]
    before_raw = snapshot(src_root)
    LOGGER.info("      policy = %s [%s]  threshold=%d  rule=%s",
                pol["policy_id"], pol["policy_status"], THRESHOLD, RULE)
    LOGGER.info("      RAW    = %s  [READ-ONLY: decoded, never written]", src_root)
    LOGGER.info("      root evidence = %s", ev["dataset_root_locator"])

    rows = read_rows(ev["paths"]["manifest"])
    integ = read_rows(ev["paths"]["integrity"])
    missing_cols = [c for c in (C_KEY, C_IMGPATH, C_MASKPATHS, C_CLASS, C_NMASK,
                                C_PAIRSTAT) if rows and c not in rows[0]]
    if missing_cols:
        LOGGER.error("      declared column(s) absent: %s", join_list(missing_cols))
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4
    keys = [r[C_KEY].strip() for r in rows]
    dupes = [k for k, n in Counter(keys).items() if n > 1]
    LOGGER.info("      manifest rows=%d  unique %s=%d  duplicates=%d",
                len(rows), C_KEY, len(set(keys)), len(dupes))
    if dupes:
        LOGGER.error("      %s is not unique; the join is ambiguous and is not "
                     "resolved silently: %s", C_KEY, dupes[:5])
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4

    # ---- 2. exclusions, read from the integrity artifact ---------------------
    LOGGER.info("[2/8] mask-integrity exclusions (from %s) ...", INTEGRITY_FILE)
    excl_map: Dict[str, dict] = {}
    for r in integ:
        k = str(r.get("sample_key", "")).strip()
        if not k:
            continue
        status = exclusion_for(str(r.get("problem_flags", "")),
                               str(r.get("mismatch_kind", "")))
        excl_map[k] = {"status": status,
                       "flags": str(r.get("problem_flags", "")),
                       "kind": str(r.get("mismatch_kind", "")),
                       "note": str(r.get("mismatch_note", ""))[:120]}
        LOGGER.info("      %-24s %-38s %s", k, status, r.get("problem_flags", ""))
    unknown = [k for k in excl_map if k not in set(keys)]
    if unknown:
        LOGGER.error("      integrity sample_key(s) absent from the manifest: %s",
                     unknown)
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4

    # ---- 3. plan --------------------------------------------------------------
    LOGGER.info("[3/8] planning canonical records ...")
    records: List[dict] = []
    for r in rows:
        key = r[C_KEY].strip()
        cls = r[C_CLASS].strip().upper()
        img_rel = r[C_IMGPATH].strip()
        mask_rel = r[C_MASKPATHS].strip()
        excl = excl_map.get(key)
        eligible = excl is None
        stem = canonical_stem(key)
        folder = "Benign" if cls == "BENIGN" else (
            "Malignant" if cls == "MALIGNANT" else "")
        nd = str(r.get(C_NEARDUP, "")).strip()
        records.append({
            "dataset": DATASET,
            "canonical_record_id": f"{DATASET}__{stem}",
            "sample_key": key,
            "source_image_id": r.get(C_IMGNAME, "").strip() or Path(img_rel).name,
            "source_mask_id": r.get(C_MASKNAMES, "").strip() or Path(mask_rel).name,
            "source_image_path": img_rel,
            "source_mask_path": mask_rel,
            "source_label": cls,
            "canonical_binary_label": (CANONICAL_LABELS.get(cls, "UNRESOLVED")
                                       if eligible else "NOT_IN_CANONICAL_COHORT"),
            "canonical_image_path": (f"{folder}/Images/{stem}.png"
                                     if eligible and folder else ""),
            "canonical_mask_path": (f"{folder}/Masks/{stem}_mask.png"
                                    if eligible and folder else ""),
            "patient_id": "", "case_id": "", "lesion_id": "",
            "classification_available": "TRUE" if eligible else "FALSE",
            "segmentation_available": "TRUE" if eligible else "FALSE",
            "image_conversion_status": "PLANNED" if eligible else "NOT_APPLICABLE",
            "mask_conversion_status": "PLANNED" if eligible else "NOT_APPLICABLE",
            "identity_provenance": (f"{MANIFEST_FILE} :: csv:[{C_PIDSTAT}]="
                                    f"{r.get(C_PIDSTAT, '').strip()}; IMAGE_ONLY"),
            "label_provenance": (f"{MANIFEST_FILE} :: csv:[{C_CLASS}] "
                                 f"(source={r.get(C_CLASSSRC, '').strip()})"),
            "mask_provenance": ("SOURCE_NATIVE_SINGLE_MASK_BINARISED_BY["
                                f"{pol['policy_id']}]" if eligible
                                else "EXCLUDED_NOT_CONVERTED"),
            "source_image_height": r.get(C_H, ""), "source_image_width": r.get(C_W, ""),
            "source_mask_height": "", "source_mask_width": "",
            "canonical_image_height": "", "canonical_image_width": "",
            "canonical_mask_height": "", "canonical_mask_width": "",
            "source_mask_encoding": r.get(C_ENC, "").strip(),
            "canonical_mask_encoding": ("SINGLE_CHANNEL_BINARY_0_255" if eligible
                                        else ""),
            "canonical_binarization_threshold": THRESHOLD if eligible else "",
            "canonical_binarization_rule": RULE if eligible else "",
            "threshold_provenance": (pol.get("threshold_provenance", "") if eligible
                                     else ""),
            "exclusion_status": ("INCLUDED_CANONICAL_JOINT" if eligible
                                 else excl["status"]),
            "exclusion_reason": ("" if eligible else
                                 f"{excl['flags']}; {excl['kind']}; {excl['note']}"),
            "notes": (f"{C_PAIRSTAT}={r.get(C_PAIRSTAT, '').strip()}; "
                      f"{C_NMASK}={r.get(C_NMASK, '').strip()}"
                      + (f"; {C_NEARDUP}={nd} (SCREENING_CANDIDATE_ONLY, not a "
                         "confirmed duplicate and not a group lock)" if nd else "")),
            "_eligible": eligible,
        })
    eligible = [r for r in records if r["_eligible"]]
    excluded = [r for r in records if not r["_eligible"]]
    by_cls = Counter(r["source_label"] for r in records)
    el_cls = Counter(r["source_label"] for r in eligible)
    LOGGER.info("      source %d (B=%d M=%d)", len(records), by_cls.get("BENIGN", 0),
                by_cls.get("MALIGNANT", 0))
    LOGGER.info("      eligible %d (B=%d M=%d)  excluded %d", len(eligible),
                el_cls.get("BENIGN", 0), el_cls.get("MALIGNANT", 0), len(excluded))
    for r in excluded:
        LOGGER.info("         EXCLUDED %-24s %s", r["sample_key"],
                    r["exclusion_status"])

    # ---- 4. resolve and pre-flight every eligible pair -----------------------
    LOGGER.info("[4/8] resolving and pre-checking source pairs ...")
    img_ok = mask_ok = plans = 0
    unresolved: List[str] = []
    for i, r in enumerate(records, 1):
        if not r["_eligible"]:
            continue
        sp = src_root / r["source_image_path"]
        mp = src_root / r["source_mask_path"]
        if not sp.is_file():
            unresolved.append(f"image missing: {r['source_image_path']}")
            r["image_conversion_status"] = "UNRESOLVED_SOURCE"
            continue
        try:
            with Image.open(sp) as im:
                r["source_image_width"] = r["canonical_image_width"] = im.width
                r["source_image_height"] = r["canonical_image_height"] = im.height
            img_ok += 1
        except Exception as exc:  # noqa: BLE001
            unresolved.append(f"image undecodable: {r['source_image_path']} "
                              f"({type(exc).__name__})")
            r["image_conversion_status"] = "UNDECODABLE_SOURCE"
            continue
        if not mp.is_file():
            unresolved.append(f"mask missing: {r['source_mask_path']}")
            r["mask_conversion_status"] = "UNRESOLVED_SOURCE"
            continue
        try:
            with Image.open(mp) as mi:
                arr = np.array(mi)
                r["source_mask_width"] = r["canonical_mask_width"] = mi.width
                r["source_mask_height"] = r["canonical_mask_height"] = mi.height
        except Exception as exc:  # noqa: BLE001
            unresolved.append(f"mask undecodable: {r['source_mask_path']} "
                              f"({type(exc).__name__})")
            r["mask_conversion_status"] = "UNDECODABLE_SOURCE"
            continue
        mask_ok += 1
        if arr.ndim == 3:
            if arr.shape[2] < 3 or not (
                    np.array_equal(arr[..., 0], arr[..., 1])
                    and np.array_equal(arr[..., 1], arr[..., 2])):
                unresolved.append(f"mask channels not replicated: "
                                  f"{r['source_mask_path']}")
                r["mask_conversion_status"] = "RGB_NOT_REPLICATED"
                continue
            gray = arr[..., 0]
        else:
            gray = arr
        if (r["source_mask_height"], r["source_mask_width"]) != (
                r["source_image_height"], r["source_image_width"]):
            unresolved.append(f"dimension mismatch not excluded: {r['sample_key']}")
            r["mask_conversion_status"] = "DIMENSION_MISMATCH"
            continue
        if int((gray >= THRESHOLD).sum()) == 0:
            unresolved.append(f"empty after threshold: {r['sample_key']}")
            r["mask_conversion_status"] = "EMPTY_AFTER_THRESHOLD"
            continue
        plans += 1
        if i % 200 == 0:
            LOGGER.info("      %d/%d", i, len(records))
    LOGGER.info("      image mappings=%d  mask mappings=%d  threshold plans=%d  "
                "unresolved=%d", img_ok, mask_ok, plans, len(unresolved))
    for u in unresolved[:10]:
        LOGGER.error("         %s", u)

    # ---- 5. execute -----------------------------------------------------------
    png_written = 0
    if mode == "execute":
        LOGGER.info("[5/8] executing canonical conversion ...")
        for r in eligible:
            if r["image_conversion_status"] != "PLANNED":
                continue
            dst = bank_ds / r["canonical_image_path"]
            dst.parent.mkdir(parents=True, exist_ok=True)
            if not dst.exists():
                with Image.open(src_root / r["source_image_path"]) as im:
                    im.save(dst, format="PNG")          # no geometry change
                png_written += 1
            r["image_conversion_status"] = "CONVERTED"
            with Image.open(src_root / r["source_mask_path"]) as mi:
                arr = np.array(mi)
            gray = arr[..., 0] if arr.ndim == 3 else arr
            canon = (255 * (gray >= THRESHOLD)).astype(np.uint8)   # the locked rule
            mdst = bank_ds / r["canonical_mask_path"]
            mdst.parent.mkdir(parents=True, exist_ok=True)
            if not mdst.exists():
                Image.fromarray(canon, mode="L").save(mdst, format="PNG")
                png_written += 1
            r["mask_conversion_status"] = "CONVERTED"
    else:
        LOGGER.info("[5/8] dry-run: no canonical PNG is written")

    # ---- 6. tests --------------------------------------------------------------
    LOGGER.info("[6/8] validation ...")
    after_bank = snapshot(prep_root / BANK_NAME)
    pre_splits = [k for k in before_bank if is_split_like(k)]
    new_splits = [k for k in after_bank if k not in before_bank and is_split_like(k)]
    mod_splits = [k for k in after_bank if k in before_bank
                  and after_bank[k] != before_bank[k] and is_split_like(k)]
    new_png = [k for k in after_bank
               if k not in before_bank and k.lower().endswith(".png")]
    raw_diff = diff_state(before_raw, snapshot(src_root))
    after_audit: Dict[str, Tuple[int, int]] = {}
    for d in prep_dirs:
        after_audit.update(snapshot(d))
    aud_diff = diff_state(before_audit, after_audit)
    paths = [r["canonical_image_path"] for r in eligible if r["canonical_image_path"]]
    paths += [r["canonical_mask_path"] for r in eligible if r["canonical_mask_path"]]
    collisions = [n for n, c in Counter(paths).items() if c > 1]

    T: List[Tuple[str, str, Optional[bool], str]] = []

    def add(i: str, d: str, ok: Optional[bool], b: str) -> None:
        T.append((i, d, ok, b))

    add("T01", f"source records == {EXPECT['records']}",
        len(records) == EXPECT["records"], f"{len(records)}")
    add("T02", f"source benign == {EXPECT['benign']}",
        by_cls.get("BENIGN") == EXPECT["benign"], f"{by_cls.get('BENIGN')}")
    add("T03", f"source malignant == {EXPECT['malignant']}",
        by_cls.get("MALIGNANT") == EXPECT["malignant"], f"{by_cls.get('MALIGNANT')}")
    add("T04", f"excluded == {EXPECT['excluded']}", len(excluded) == EXPECT["excluded"],
        join_list([r["sample_key"] for r in excluded], ", "))
    add("T05", f"eligible == {EXPECT['eligible']}",
        len(eligible) == EXPECT["eligible"], f"{len(eligible)}")
    add("T06", f"eligible benign == {EXPECT['eligible_benign']}",
        el_cls.get("BENIGN") == EXPECT["eligible_benign"], f"{el_cls.get('BENIGN')}")
    add("T07", f"eligible malignant == {EXPECT['eligible_malignant']}",
        el_cls.get("MALIGNANT") == EXPECT["eligible_malignant"],
        f"{el_cls.get('MALIGNANT')}")
    add("T08", "every excluded record has an explicit status and no canonical path",
        all(r["exclusion_status"].startswith("EXCLUDE") and not r["canonical_image_path"]
            and not r["canonical_mask_path"] for r in excluded),
        join_list([f"{r['sample_key']}={r['exclusion_status']}" for r in excluded],
                  "; "))
    add("T09", f"source image mappings == {EXPECT['eligible']}",
        img_ok == EXPECT["eligible"], f"{img_ok}/{len(eligible)}")
    add("T10", f"source mask mappings == {EXPECT['eligible']}",
        mask_ok == EXPECT["eligible"], f"{mask_ok}/{len(eligible)}")
    add("T11", f"canonical threshold plans == {EXPECT['eligible']}",
        plans == EXPECT["eligible"], f"{plans} at threshold {THRESHOLD}")
    add("T12", "threshold comes from the locked policy, not from this file",
        THRESHOLD == int(pol["canonical_threshold"])
        and pol.get("threshold_provenance") == "PHASE2B_CANONICAL_POLICY",
        f"{THRESHOLD} @ {POLICY_FILE}::json:canonical_threshold")
    add("T13", "union plans == 0 (BUS-UC has no multi-mask case)",
        all(str(r.get(C_NMASK, "1")).strip() == "1" for r in rows),
        f"mask_count values: {dict(Counter(str(r.get(C_NMASK, '')).strip() for r in rows))}")
    add("T14", "no patient, case or lesion identity invented",
        all(not r["patient_id"] and not r["case_id"] and not r["lesion_id"]
            for r in records), "all three empty on every record")
    add("T15", "no duplicate group or split lock invented from screening candidates",
        all(not str(r.get(C_EXACTDUP, "")).strip() for r in rows)
        and "duplicate_group_id" not in SCHEMA_FIELDS,
        f"{sum(1 for r in rows if str(r.get(C_NEARDUP, '')).strip())} screening "
        "candidate flag(s) preserved in notes only")
    add("T16", "no unresolved eligible record", not unresolved, f"{len(unresolved)}")
    add("T17", "no canonical path collision", not collisions, f"{len(collisions)}")
    add("T18", "RAW unchanged", raw_diff["unchanged"],
        f"modified {len(raw_diff['modified'])}, added {len(raw_diff['added'])}")
    add("T19", "completed audit artifacts unchanged", aud_diff["unchanged"],
        f"modified {len(aud_diff['modified'])}")
    add("T20", "no split created or modified by this run",
        not new_splits and not mod_splits,
        f"preexisting={len(pre_splits)}, new={len(new_splits)}, "
        f"modified={len(mod_splits)}")
    add("T21", ("dry-run wrote no canonical PNG" if mode == "dry-run"
                else f"execution wrote {EXPECT['png']} canonical PNG"),
        (not new_png) if mode == "dry-run" else (len(new_png) == EXPECT["png"]),
        f"{len(new_png)} new PNG")
    add("T22", "no split field in the manifest schema",
        not any(is_split_like(f) for f in SCHEMA_FIELDS), "schema checked")
    add("T23", "manifest retains all source records",
        len(records) == EXPECT["records"]
        and len(eligible) + len(excluded) == len(records),
        f"{len(eligible)} + {len(excluded)} = {len(records)}")
    add("T24", "no preprocessing in the conversion path", None,
        "code-scope declaration: decode, channel reduction, threshold, lossless PNG; "
        "no resize, resample, crop, pad, CLAHE, normalisation or augmentation call")
    verified = [t for t in T if t[2] is not None]
    failed = [t for t in verified if not t[2]]
    for t in T:
        (LOGGER.info if t[2] is not False else LOGGER.error)(
            "      %-5s [%-11s] %-52s %s", t[0],
            "DECLARATION" if t[2] is None else ("PASS" if t[2] else "FAIL"),
            t[1][:52], t[3][:50])

    # ---- 7. outputs -------------------------------------------------------------
    LOGGER.info("[7/8] writing outputs ...")
    tag = "DRY_RUN_" if mode == "dry-run" else ""
    out_rows = [{k: r[k] for k in SCHEMA_FIELDS} for r in records]
    write_csv(manifests / f"{tag}canonical_manifest.csv", SCHEMA_FIELDS, out_rows)
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "canonical_manifest"
        ws.append(SCHEMA_FIELDS)
        for r in out_rows:
            ws.append([r.get(h, "") for h in SCHEMA_FIELDS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = (f"A1:{get_column_letter(len(SCHEMA_FIELDS))}"
                              f"{len(out_rows) + 1}")
        wb.save(manifests / f"{tag}canonical_manifest.xlsx")
        LOGGER.info("WROTE %-48s (1 sheet)", f"{tag}canonical_manifest.xlsx")
    except ImportError:
        LOGGER.warning("openpyxl not installed; manifest XLSX skipped")

    blocked = bool(failed)
    state = {
        "phase": "2B.3.2", "dataset": DATASET, "mode": mode, "generated": now_iso(),
        "builder_sha256": sha256_file(Path(__file__).resolve()),
        "schema_sha256": sha256_text(join_list(SCHEMA_FIELDS, ",")),
        "evidence_hashes": ev["hashes"],
        "source_root": str(src_root),
        "policy": {"policy_id": pol["policy_id"],
                   "policy_status": pol["policy_status"],
                   "canonical_threshold": THRESHOLD,
                   "threshold_provenance": pol.get("threshold_provenance")},
        "counts": {"source_records": len(records),
                   "source_benign": by_cls.get("BENIGN", 0),
                   "source_malignant": by_cls.get("MALIGNANT", 0),
                   "eligible": len(eligible),
                   "eligible_benign": el_cls.get("BENIGN", 0),
                   "eligible_malignant": el_cls.get("MALIGNANT", 0),
                   "excluded": len(excluded),
                   "threshold_plans": plans, "union_plans": 0,
                   "confirmed_duplicate_groups": 0, "group_locks": 0,
                   "near_duplicate_candidates_flagged": sum(
                       1 for r in rows if str(r.get(C_NEARDUP, "")).strip())},
        "excluded_records": [{"sample_key": r["sample_key"],
                              "status": r["exclusion_status"],
                              "reason": r["exclusion_reason"]} for r in excluded],
        "image_mappings_resolved": img_ok, "mask_mappings_resolved": mask_ok,
        "unresolved": unresolved[:50], "collisions": collisions,
        "canonical_png_written": png_written,
        "preexisting_split_artifacts": len(pre_splits),
        "newly_created_split_artifacts": len(new_splits),
        "modified_split_artifacts": len(mod_splits),
        "raw_modified": not raw_diff["unchanged"],
        "audit_modified": not aud_diff["unchanged"],
        "tests_pass": sum(1 for t in verified if t[2]), "tests_fail": len(failed),
        "verdict": ("BLOCKED" if blocked else
                    "READY_FOR_BUS_UC_CANONICAL_EXECUTION" if mode == "dry-run"
                    else "BUS_UC_CANONICAL_BUILD_COMPLETE"),
    }
    write_text(reports / f"{tag}canonical_build_report.json",
               json.dumps(state, indent=2, ensure_ascii=False, default=str))

    # ---- 8. report ----------------------------------------------------------------
    LOGGER.info("[8/8] report")
    LOGGER.info("=" * 70)
    LOGGER.info("PHASE 2B.3.2 BUS-UC CANONICAL BANK - %s", mode.upper())
    LOGGER.info("")
    LOGGER.info("Evidence:")
    LOGGER.info("  manifest:   %s (%d rows, key=%s)", MANIFEST_FILE, len(rows), C_KEY)
    LOGGER.info("  integrity:  %s (%d rows)", INTEGRITY_FILE, len(integ))
    LOGGER.info("  policy:     %s [%s] threshold=%d (%s)", pol["policy_id"],
                pol["policy_status"], THRESHOLD, pol.get("threshold_provenance"))
    LOGGER.info("")
    LOGGER.info("Source:     %d  (benign %d / malignant %d)", len(records),
                by_cls.get("BENIGN", 0), by_cls.get("MALIGNANT", 0))
    LOGGER.info("Eligible:   %d  (benign %d / malignant %d)", len(eligible),
                el_cls.get("BENIGN", 0), el_cls.get("MALIGNANT", 0))
    LOGGER.info("Excluded:   %d", len(excluded))
    for r in excluded:
        LOGGER.info("    %-24s %s", r["sample_key"], r["exclusion_status"])
    LOGGER.info("")
    LOGGER.info("source image mappings:   %d", img_ok)
    LOGGER.info("source mask mappings:    %d", mask_ok)
    LOGGER.info("threshold plans:         %d at %d", plans, THRESHOLD)
    LOGGER.info("union plans:             0")
    LOGGER.info("patient/case/lesion IDs invented: 0 / 0 / 0")
    LOGGER.info("confirmed duplicate groups: 0    group locks: 0")
    LOGGER.info("near-duplicate candidates flagged: %d (screening only)",
                state["counts"]["near_duplicate_candidates_flagged"])
    LOGGER.info("unresolved eligible:     %d", len(unresolved))
    LOGGER.info("filename collisions:     %d", len(collisions))
    LOGGER.info("")
    LOGGER.info("RAW modified:            %s", state["raw_modified"])
    LOGGER.info("audit modified:          %s", state["audit_modified"])
    LOGGER.info("canonical PNG written:   %d", png_written)
    LOGGER.info("splits created/modified: %d / %d", len(new_splits), len(mod_splits))
    LOGGER.info("")
    LOGGER.info("TESTS PASS: %d", state["tests_pass"])
    LOGGER.info("TESTS FAIL: %d", state["tests_fail"])
    LOGGER.info("")
    LOGGER.info("FINAL VERDICT:")
    if blocked:
        LOGGER.error("BLOCKED")
        for t in failed:
            LOGGER.error("    FAILED %s: %s | %s", t[0], t[1], t[3])
        return 4
    if mode == "dry-run":
        LOGGER.info("READY_FOR_BUS_UC_CANONICAL_EXECUTION")
        LOGGER.info("=" * 70)
        LOGGER.info("STOP. --execute is not run automatically and needs human "
                    "approval.")
    else:
        LOGGER.info("BUS_UC_CANONICAL_BUILD_COMPLETE")
    return 0


def execute_gate(prep_root: Path) -> Tuple[bool, str]:
    st = (prep_root / BANK_NAME / DATASET / "reports"
          / "DRY_RUN_canonical_build_report.json")
    if not st.is_file():
        return False, f"no dry-run report at {st}"
    d = json.loads(st.read_text(encoding="utf-8"))
    if d.get("verdict") != "READY_FOR_BUS_UC_CANONICAL_EXECUTION":
        return False, f"the last dry-run verdict was {d.get('verdict')!r}"
    if d.get("builder_sha256") != sha256_file(Path(__file__).resolve()):
        return False, "the builder changed since the approved dry-run; rerun --dry-run"
    if d.get("schema_sha256") != sha256_text(join_list(SCHEMA_FIELDS, ",")):
        return False, "the schema changed since the approved dry-run; rerun --dry-run"
    ev = load_evidence(prep_root)
    for name, h in (d.get("evidence_hashes") or {}).items():
        if ev["hashes"].get(name) != h:
            return False, f"evidence artifact {name} changed since the dry-run"
    return True, "the dry-run approval matches the current builder, schema and evidence"


def main(argv: Optional[Sequence[str]] = None) -> int:
    p = argparse.ArgumentParser(description="BUS-UC canonical bank builder.")
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--dry-run", action="store_true")
    g.add_argument("--execute", action="store_true")
    p.add_argument("--prep-root", default=DEFAULT_PREP_ROOT)
    a = p.parse_args(argv)
    mode = "dry-run" if a.dry_run else "execute"
    if mode == "execute":
        ok, why = execute_gate(Path(a.prep_root))
        if not ok:
            sys.stderr.write(f"EXECUTE refused: {why}\n")
            return 5
        sys.stderr.write(f"execute gate: {why}\n")
    try:
        return run(mode, Path(a.prep_root))
    except Exception:  # noqa: BLE001
        if LOGGER.handlers:
            LOGGER.critical("UNHANDLED EXCEPTION\n%s", traceback.format_exc())
        else:
            sys.stderr.write(traceback.format_exc())
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
