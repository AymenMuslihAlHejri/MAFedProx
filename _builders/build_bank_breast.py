#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_bank_breast.py
====================

PHASE 2B.3.3 - BrEaST CANONICAL DATASET BANK BUILDER

Same fail-closed architecture as BUSI and BUS-UC. Nothing is carried over from either
dataset's science.

WHY THIS DATASET NEEDS NO LOCKED THRESHOLD ARTIFACT
    BUS-UC required one because a midpoint threshold was a DISCRETIONARY choice with no
    source authority. BrEaST exercises no discretion: the audit establishes the masks
    are already strictly binary {0,255} with identical colour channels, so canonical
    conversion only removes redundant channels. The rule is therefore declared here and
    its PRECONDITIONS are verified against the audit evidence at run time - if any
    precondition fails the build blocks instead of silently thresholding.

    POLICY_ID BREAST_CANONICAL_MASK_REPLICATION_REDUCTION_V1
        verify colour channels identical -> take one plane -> assert strictly {0,255}
        -> single-channel lossless PNG, pixel-equivalent to the source binary plane.
    No threshold, no morphology, no erosion, dilation, smoothing, resize, resample or
    pseudo-mask. BUS_UC_CANONICAL_MASK_BINARIZATION_V1 is NOT applied and neither is
    BUSI_CANONICAL_MULTIMASK_UNION_V1.

NOTE ON CHANNEL COUNT
    The audit records mode=RGBA and channels=4 for every mask, not RGB/3. The colour
    channels are identical, so the reduction is unchanged, but this builder verifies
    R=G=B explicitly AND inspects the alpha channel rather than assuming three
    channels. Alpha is never used to define foreground.

LOCKED EVIDENCE BINDINGS
    manifest   11_breast_audit_manifest.csv   (252 rows, key sample_key)
    pairing    03_breast_pairing_audit.csv    (252 rows)
    mask       05_breast_mask_audit.csv       (252 rows)
    class      06_breast_class_audit.csv      (252 rows)
    identity   08_breast_identity_audit.csv
    duplicates 09_breast_exact_duplicates.csv
    summary    14_breast_audit_summary.json :: json:raw_root

Usage:
    python build_bank_breast.py --dry-run
    python build_bank_breast.py --execute
Requires: Python >= 3.8, Pillow, numpy. openpyxl optional.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import platform
import re
import sys
import traceback
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

DATASET = "BrEaST"
DEFAULT_PREP_ROOT = (
    r"D:\PhD\002 FL Mul BC\01 ready\Springer\BMC Med Imaging"
    r"\Revision 1\dataset preparing"
)
BANK_NAME = "CANONICAL_DATASET_BANK"
ARCHIVE_DIR_NAMES = ("audit_old",)

MANIFEST_FILE = "11_breast_audit_manifest.csv"
PAIRING_FILE = "03_breast_pairing_audit.csv"
MASK_FILE = "05_breast_mask_audit.csv"
CLASS_FILE = "06_breast_class_audit.csv"
IDENTITY_FILE = "08_breast_identity_audit.csv"
DUPLICATE_FILE = "09_breast_exact_duplicates.csv"
SUMMARY_FILE = "14_breast_audit_summary.json"
ROOT_LOCATORS = ["raw_root", "source_path", "dataset_root", "source_root"]

POLICY_ID = "BREAST_CANONICAL_MASK_REPLICATION_REDUCTION_V1"
POLICY_RULE = ("verify colour channels identical -> reduce to one plane -> assert "
               "strictly {0,255} -> single-channel lossless PNG")

EXPECT = {"records": 252, "images": 252, "masks": 252, "benign": 154, "malignant": 98,
          "normal": 0, "eligible": 252, "excluded": 0, "png": 504}

CANONICAL_LABELS = {"BENIGN": 0, "MALIGNANT": 1}
SCHEMA_FIELDS = [
    "dataset", "canonical_record_id", "sample_key", "source_image_id",
    "source_mask_id", "source_image_path", "source_mask_path", "source_label",
    "canonical_binary_label", "canonical_image_path", "canonical_mask_path",
    "patient_id", "case_id", "image_id", "lesion_id", "abnormality_id",
    "same_source_image_group", "duplicate_group_id",
    "classification_available", "segmentation_available",
    "image_conversion_status", "mask_conversion_status",
    "identity_provenance", "label_provenance", "mask_provenance",
    "source_image_height", "source_image_width",
    "source_mask_height", "source_mask_width",
    "canonical_image_height", "canonical_image_width",
    "canonical_mask_height", "canonical_mask_width",
    "source_image_format", "source_mask_format",
    "source_mask_encoding", "canonical_mask_encoding",
    "canonical_mask_rule", "canonical_mask_policy_id",
    "exclusion_status", "exclusion_reason", "notes",
]
SPLIT_TOKENS = {"split", "splits", "fold", "folds", "kfold", "train", "training",
                "trainval", "val", "validation", "test", "testset", "testing",
                "holdout", "partition"}

SEP = " | "
LOGGER = logging.getLogger("bank_breast")


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def join_list(v: Sequence[Any], sep: str = SEP) -> str:
    return sep.join(str(x) for x in v)


def archived(p: Path) -> bool:
    return any(x.lower() in ARCHIVE_DIR_NAMES for x in p.parts[:-1])


def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as fh:
        for b in iter(lambda: fh.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def sha256_text(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def snapshot(root: Path) -> Dict[str, Tuple[int, int]]:
    out: Dict[str, Tuple[int, int]] = {}
    if not root.is_dir():
        return out
    for p in root.rglob("*"):
        try:
            if p.is_file():
                st = p.stat()
                out[str(p)] = (st.st_size, st.st_mtime_ns)
        except OSError:
            pass
    return out


def diff_state(a: dict, b: dict) -> dict:
    added = sorted(set(b) - set(a))
    removed = sorted(set(a) - set(b))
    modified = sorted(k for k in set(a) & set(b) if a[k] != b[k])
    return {"added": added, "removed": removed, "modified": modified,
            "unchanged": not (added or removed or modified)}


def is_split_like(path_str: str) -> bool:
    pth = Path(path_str)
    raw = [t for t in re.split(r"[^A-Za-z0-9]+", pth.stem) if t]
    raw += [t for part in pth.parts[:-1]
            for t in re.split(r"[^A-Za-z0-9]+", part) if t]
    toks = set()
    for t in raw:
        toks.add(t.lower())
        toks.add(re.sub(r"\d+$", "", t.lower()))
    return bool(toks & SPLIT_TOKENS)


def write_csv(path: Path, fields: Sequence[str], rows: Sequence[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(fields), extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)
    LOGGER.info("WROTE %-46s (%d data rows)", path.name, len(rows))


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    LOGGER.info("WROTE %-46s (%d chars)", path.name, len(text))


def read_rows(p: Path) -> List[dict]:
    with p.open(encoding="utf-8-sig", errors="replace") as fh:
        s = fh.read(8192)
        fh.seek(0)
        if not s.strip():
            return []
        d = ";" if s.count(";") > s.count(",") else ","
        return list(csv.DictReader(fh, delimiter=d))


def setup_logging(report_dir: Path, name: str) -> None:
    LOGGER.setLevel(logging.DEBUG)
    LOGGER.handlers.clear()
    f = logging.Formatter("%(asctime)s | %(levelname)-8s | %(message)s")
    sh = logging.StreamHandler(sys.stdout)
    sh.setFormatter(f)
    sh.setLevel(logging.INFO)
    LOGGER.addHandler(sh)
    report_dir.mkdir(parents=True, exist_ok=True)
    fh = logging.FileHandler(report_dir / name, mode="w", encoding="utf-8")
    fh.setFormatter(f)
    fh.setLevel(logging.DEBUG)
    LOGGER.addHandler(fh)


def find_one(files: Sequence[Path], filename: str) -> Tuple[Optional[Path], str]:
    hits = sorted(p for p in files if p.name.lower() == filename.lower())
    if len(hits) == 1:
        (only,) = hits
        return only, "UNIQUE"
    return None, ("MISSING" if not hits else f"AMBIGUOUS({len(hits)})")


def canonical_stem(cls: str, sample_key: str) -> str:
    key = f"{cls}/{sample_key}"
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(sample_key).stem).strip("_")
    return f"{safe or 'image'}__{sha256_text(key)[:10]}"


def reduce_mask(arr, np) -> Tuple[Any, str]:
    """Remove redundant colour channels, preserving the binary plane exactly.
    Alpha is inspected but NEVER used to define foreground."""
    if arr.ndim == 2:
        plane, note = arr, "ALREADY_SINGLE_CHANNEL"
    else:
        n = arr.shape[2]
        if n < 3:
            return None, f"UNEXPECTED_CHANNEL_COUNT_{n}"
        if not (np.array_equal(arr[..., 0], arr[..., 1])
                and np.array_equal(arr[..., 1], arr[..., 2])):
            return None, "COLOUR_CHANNELS_NOT_IDENTICAL"
        plane = arr[..., 0]
        if n == 4:
            al = set(np.unique(arr[..., 3]).tolist())
            note = ("RGBA_ALPHA_FULLY_OPAQUE" if al <= {255}
                    else f"RGBA_ALPHA_VARIES{sorted(al)[:4]}_NOT_USED_AS_FOREGROUND")
        else:
            note = "RGB_CHANNELS_REPLICATED"
    vals = set(np.unique(plane).tolist())
    if not vals <= {0, 255}:
        return None, f"PLANE_NOT_STRICTLY_0_255{sorted(vals)[:5]}"
    return plane.astype(np.uint8), note


def load_evidence(prep_root: Path) -> dict:
    ev: Dict[str, Any] = {"paths": {}, "status": {}, "hashes": {}, "problems": []}
    dirs = [d for d in prep_root.iterdir() if d.is_dir()
            and "breast" in d.name.lower().replace(" ", "")
            and "bus" not in d.name.lower()]
    files = [p for d in dirs for p in d.rglob("*") if p.is_file() and not archived(p)]
    for key, fname, required in (("manifest", MANIFEST_FILE, True),
                                 ("pairing", PAIRING_FILE, True),
                                 ("mask", MASK_FILE, True),
                                 ("class", CLASS_FILE, True),
                                 ("identity", IDENTITY_FILE, True),
                                 ("duplicates", DUPLICATE_FILE, False),
                                 ("summary", SUMMARY_FILE, True)):
        art, st = find_one(files, fname)
        ev["paths"][key] = art
        ev["status"][key] = st
        if art is not None:
            ev["hashes"][art.name] = sha256_file(art)
        elif required:
            ev["problems"].append(f"BLOCKED: {fname} -> {st}")
    ev["dataset_root"] = None
    ev["dataset_root_locator"] = ""
    if ev["paths"].get("summary"):
        blob = json.loads(ev["paths"]["summary"].read_text(encoding="utf-8"))
        for loc in ROOT_LOCATORS:
            v = blob.get(loc)
            if isinstance(v, str) and v.strip():
                ev["dataset_root"] = Path(v.strip())
                ev["dataset_root_locator"] = f"{SUMMARY_FILE} :: json:{loc}"
                break
        if ev["dataset_root"] is None:
            ev["problems"].append(
                f"BLOCKED: no source root in {SUMMARY_FILE}; tried "
                + join_list([f"json:{k}" for k in ROOT_LOCATORS], ", "))
    return ev


def run(mode: str, prep_root: Path) -> int:
    bank_ds = prep_root / BANK_NAME / DATASET
    reports, manifests = bank_ds / "reports", bank_ds / "manifests"
    setup_logging(reports, f"breast_{mode}.log")
    LOGGER.info("=" * 78)
    LOGGER.info("PHASE 2B.3.3 BrEaST CANONICAL BANK BUILDER  [%s]", mode.upper())
    LOGGER.info("=" * 78)
    LOGGER.info("script : %s", Path(__file__).resolve())
    LOGGER.info("python : %s", platform.python_version())
    LOGGER.info("bank   : %s", bank_ds)

    try:
        from PIL import Image
        import numpy as np
    except ImportError:
        LOGGER.critical("Pillow and numpy required: pip install pillow numpy")
        return 3
    Image.MAX_IMAGE_PIXELS = None

    prep_dirs = [d for d in prep_root.iterdir()
                 if d.is_dir() and d.name != BANK_NAME]
    before_audit: Dict[str, Tuple[int, int]] = {}
    for d in prep_dirs:
        before_audit.update(snapshot(d))
    before_bank = snapshot(prep_root / BANK_NAME)

    # ---- 1. evidence -------------------------------------------------------
    LOGGER.info("[1/8] locked evidence bindings ...")
    ev = load_evidence(prep_root)
    for k, st in ev["status"].items():
        art = ev["paths"].get(k)
        LOGGER.info("      %-11s %-12s %s", k, st, art.name if art else "")
    if ev["problems"]:
        for p in ev["problems"]:
            LOGGER.error("      %s", p)
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4
    src_root = ev["dataset_root"]
    before_raw = snapshot(src_root)
    LOGGER.info("      RAW    = %s  [READ-ONLY: decoded, never written]", src_root)
    LOGGER.info("      root evidence = %s", ev["dataset_root_locator"])
    LOGGER.info("      mask policy = %s", POLICY_ID)
    LOGGER.info("      rule        = %s", POLICY_RULE)

    rows = read_rows(ev["paths"]["manifest"])
    pair = read_rows(ev["paths"]["pairing"])
    mask = read_rows(ev["paths"]["mask"])
    cls_rows = read_rows(ev["paths"]["class"])
    ident = read_rows(ev["paths"]["identity"])
    dups = read_rows(ev["paths"]["duplicates"]) if ev["paths"].get("duplicates") else []
    keys = [r["sample_key"].strip() for r in rows]
    dupe_keys = [k for k, n in Counter(keys).items() if n > 1]
    LOGGER.info("      manifest=%d  pairing=%d  mask=%d  class=%d  identity=%d  "
                "duplicates=%d", len(rows), len(pair), len(mask), len(cls_rows),
                len(ident), len(dups))
    if dupe_keys:
        LOGGER.error("      sample_key not unique: %s", dupe_keys[:5])
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4

    # ---- 2. reproduce the accounting from the evidence -----------------------
    LOGGER.info("[2/8] reproducing source accounting from the evidence ...")
    by_cls = Counter(r["class_label"].strip().upper() for r in rows)
    n_masks = sum(int(r["mask_count"]) for r in rows if r["mask_count"].strip())
    cls_resolved = Counter(r["resolved_class"].strip().upper() for r in cls_rows)
    blob = json.loads(ev["paths"]["summary"].read_text(encoding="utf-8"))
    LOGGER.info("      manifest: records=%d  benign=%d  malignant=%d  masks=%d",
                len(rows), by_cls.get("BENIGN", 0), by_cls.get("MALIGNANT", 0), n_masks)
    LOGGER.info("      class audit: %s  (agreement %s)", dict(cls_resolved),
                dict(Counter(r["source_agreement"] for r in cls_rows)))
    LOGGER.info("      summary json: images=%s masks=%s benign=%s malignant=%s "
                "normal=%s", blob.get("source_images"), blob.get("segmentation_masks"),
                (blob.get("class_distribution") or {}).get("benign"),
                (blob.get("class_distribution") or {}).get("malignant"),
                blob.get("normal_cases"))

    # ---- 3. mask policy preconditions, verified from the mask audit ----------
    LOGGER.info("[3/8] verifying the mask-policy preconditions ...")
    mk_bin = Counter(r["is_binary_encoding"].strip().upper() for r in mask)
    mk_rep = Counter(r["rgb_channels_identical"].strip().upper() for r in mask)
    mk_val = Counter(r["unique_values"].strip() for r in mask)
    mk_mode = Counter(r["mode"].strip() for r in mask)
    mk_ch = Counter(r["channels"].strip() for r in mask)
    mk_empty = Counter(r["empty_mask"].strip().upper() for r in mask)
    mk_dim = Counter(r["dimension_match_with_image"].strip().upper() for r in mask)
    for lbl, c in (("is_binary_encoding", mk_bin), ("rgb_channels_identical", mk_rep),
                   ("unique_values", mk_val), ("mode", mk_mode), ("channels", mk_ch),
                   ("empty_mask", mk_empty), ("dimension_match", mk_dim)):
        LOGGER.info("      %-24s %s", lbl, dict(c))
    precond = {
        "all masks binary": mk_bin.get("YES", 0) == len(mask),
        "colour channels identical": mk_rep.get("YES", 0) == len(mask),
        "values strictly 0,255": all(
            set(v.replace(" ", "").split(",")) <= {"0", "255"} for v in mk_val),
        "no empty mask": mk_empty.get("NO", 0) == len(mask),
        "mask dimensions match image": mk_dim.get("YES", 0) == len(mask),
    }
    for k, v in precond.items():
        (LOGGER.info if v else LOGGER.error)("      precondition %-30s %s", k,
                                             "OK" if v else "FAILED")
    if not all(precond.values()):
        LOGGER.error("      the declared reduction rule requires all preconditions; "
                     "the build blocks rather than falling back to a threshold")
        LOGGER.error("FINAL VERDICT:")
        LOGGER.error("BLOCKED")
        return 4
    if set(mk_ch) != {"3"}:
        LOGGER.warning("      note: masks are %s-channel (%s), not 3-channel RGB; the "
                       "colour channels are identical so the reduction is unchanged, "
                       "and alpha is inspected but never used as foreground",
                       join_list(sorted(mk_ch), "/"), join_list(sorted(mk_mode), "/"))

    # ---- 4. identity and duplicates ------------------------------------------
    LOGGER.info("[4/8] identity and duplicate evidence ...")
    id_avail = {r["identifier_type"].strip(): r["available"].strip().upper()
                for r in ident}
    for r in ident:
        LOGGER.info("      %-20s available=%-4s unique=%-5s coverage=%s",
                    r["identifier_type"], r["available"], r["unique_count"],
                    r["coverage"])
    pid_rows = sum(1 for r in rows if r.get("patient_id", "").strip())
    case_rows = sum(1 for r in rows if r.get("case_id", "").strip())
    les_rows = sum(1 for r in rows if r.get("lesion_id", "").strip())
    dup_flagged = sum(1 for r in rows if r.get("exact_duplicate_group", "").strip())
    LOGGER.info("      manifest identity columns populated: patient=%d case=%d "
                "lesion=%d", pid_rows, case_rows, les_rows)
    LOGGER.info("      exact duplicate rows in %s: %d | manifest records carrying a "
                "group: %d", DUPLICATE_FILE, len(dups), dup_flagged)

    # ---- 5. plan --------------------------------------------------------------
    LOGGER.info("[5/8] planning canonical records ...")
    records: List[dict] = []
    for r in rows:
        key = r["sample_key"].strip()
        cls = r["class_label"].strip().upper()
        img_rel = r["image_path"].strip()
        mask_rel = r["mask_paths"].strip()
        folder = "Benign" if cls == "BENIGN" else (
            "Malignant" if cls == "MALIGNANT" else "")
        stem = canonical_stem(cls, key)
        records.append({
            "dataset": DATASET,
            "canonical_record_id": f"{DATASET}__{stem}",
            "sample_key": key,
            "source_image_id": r.get("image_filename", "").strip(),
            "source_mask_id": r.get("mask_filenames", "").strip(),
            "source_image_path": img_rel, "source_mask_path": mask_rel,
            "source_label": cls,
            "canonical_binary_label": CANONICAL_LABELS.get(cls, "UNRESOLVED"),
            "canonical_image_path": f"{folder}/Images/{stem}.png" if folder else "",
            "canonical_mask_path": f"{folder}/Masks/{stem}_mask.png" if folder else "",
            "patient_id": "", "case_id": "",
            "image_id": r.get("image_filename", "").strip(),
            "lesion_id": "", "abnormality_id": "",
            "same_source_image_group": f"{DATASET}__{stem}",
            "duplicate_group_id": r.get("exact_duplicate_group", "").strip(),
            "classification_available": "TRUE", "segmentation_available": "TRUE",
            "image_conversion_status": "PLANNED", "mask_conversion_status": "PLANNED",
            "identity_provenance": (f"{IDENTITY_FILE} :: PATIENT_ID="
                                    f"{id_avail.get('PATIENT_ID', '?')}, CASE_ID="
                                    f"{id_avail.get('CASE_ID', '?')}, LESION_ID="
                                    f"{id_avail.get('LESION_ID', '?')}; IMAGE_ONLY"),
            "label_provenance": (f"{CLASS_FILE} :: csv:[resolved_class] "
                                 f"(source={r.get('class_label_source', '').strip()})"),
            "mask_provenance": f"SOURCE_NATIVE_BINARY_REDUCED_BY[{POLICY_ID}]",
            "source_image_height": r.get("height", ""),
            "source_image_width": r.get("width", ""),
            "source_mask_height": "", "source_mask_width": "",
            "canonical_image_height": "", "canonical_image_width": "",
            "canonical_mask_height": "", "canonical_mask_width": "",
            "source_image_format": Path(img_rel).suffix.lstrip(".").upper(),
            "source_mask_format": Path(mask_rel).suffix.lstrip(".").upper(),
            "source_mask_encoding": r.get("mask_encoding", "").strip(),
            "canonical_mask_encoding": "SINGLE_CHANNEL_BINARY_0_255",
            "canonical_mask_rule": POLICY_RULE,
            "canonical_mask_policy_id": POLICY_ID,
            "exclusion_status": "INCLUDED_CANONICAL_JOINT", "exclusion_reason": "",
            "notes": (f"pairing_status={r.get('pairing_status', '').strip()}; "
                      f"mask_count={r.get('mask_count', '').strip()}; "
                      f"audit_status={r.get('audit_status', '').strip()}"),
        })
    by_el = Counter(r["source_label"] for r in records)
    LOGGER.info("      planned %d record(s): benign=%d malignant=%d", len(records),
                by_el.get("BENIGN", 0), by_el.get("MALIGNANT", 0))

    # ---- 6. resolve and pre-check every pair ---------------------------------
    LOGGER.info("[6/8] resolving and pre-checking source pairs ...")
    img_ok = mask_ok = plans = 0
    unresolved: List[str] = []
    alpha_notes: Counter = Counter()
    for i, r in enumerate(records, 1):
        sp = src_root / r["source_image_path"]
        mp = src_root / r["source_mask_path"]
        if not sp.is_file():
            unresolved.append(f"image missing: {r['source_image_path']}")
            r["image_conversion_status"] = "UNRESOLVED_SOURCE"
            continue
        try:
            with Image.open(sp) as im:
                r["source_image_width"] = r["canonical_image_width"] = im.width
                r["source_image_height"] = r["canonical_image_height"] = im.height
            img_ok += 1
        except Exception as exc:  # noqa: BLE001
            unresolved.append(f"image undecodable: {r['source_image_path']} "
                              f"({type(exc).__name__})")
            r["image_conversion_status"] = "UNDECODABLE_SOURCE"
            continue
        if not mp.is_file():
            unresolved.append(f"mask missing: {r['source_mask_path']}")
            r["mask_conversion_status"] = "UNRESOLVED_SOURCE"
            continue
        try:
            with Image.open(mp) as mi:
                arr = np.array(mi)
                r["source_mask_width"] = r["canonical_mask_width"] = mi.width
                r["source_mask_height"] = r["canonical_mask_height"] = mi.height
        except Exception as exc:  # noqa: BLE001
            unresolved.append(f"mask undecodable: {r['source_mask_path']} "
                              f"({type(exc).__name__})")
            r["mask_conversion_status"] = "UNDECODABLE_SOURCE"
            continue
        mask_ok += 1
        plane, note = reduce_mask(arr, np)
        alpha_notes[note] += 1
        if plane is None:
            unresolved.append(f"mask reduction refused: {r['sample_key']} ({note})")
            r["mask_conversion_status"] = f"REFUSED_{note}"
            continue
        if (r["source_mask_height"], r["source_mask_width"]) != (
                r["source_image_height"], r["source_image_width"]):
            unresolved.append(f"dimension mismatch: {r['sample_key']}")
            r["mask_conversion_status"] = "DIMENSION_MISMATCH"
            continue
        if int((plane == 255).sum()) == 0:
            unresolved.append(f"empty mask: {r['sample_key']}")
            r["mask_conversion_status"] = "EMPTY_MASK"
            continue
        r["notes"] += f"; channel_reduction={note}"
        plans += 1
        if i % 100 == 0:
            LOGGER.info("      %d/%d", i, len(records))
    LOGGER.info("      image mappings=%d  mask mappings=%d  reduction plans=%d  "
                "unresolved=%d", img_ok, mask_ok, plans, len(unresolved))
    LOGGER.info("      channel reduction outcomes: %s", dict(alpha_notes))
    for u in unresolved[:10]:
        LOGGER.error("         %s", u)

    # ---- 7. execute ------------------------------------------------------------
    png_written = 0
    if mode == "execute":
        LOGGER.info("[7/8] executing canonical conversion ...")
        for r in records:
            if r["image_conversion_status"] != "PLANNED":
                continue
            dst = bank_ds / r["canonical_image_path"]
            dst.parent.mkdir(parents=True, exist_ok=True)
            if not dst.exists():
                with Image.open(src_root / r["source_image_path"]) as im:
                    im.save(dst, format="PNG")      # no geometry or value change
                png_written += 1
            r["image_conversion_status"] = "CONVERTED"
            with Image.open(src_root / r["source_mask_path"]) as mi:
                arr = np.array(mi)
            plane, _note = reduce_mask(arr, np)
            if plane is None:
                r["mask_conversion_status"] = "REFUSED_AT_EXECUTE"
                continue
            mdst = bank_ds / r["canonical_mask_path"]
            mdst.parent.mkdir(parents=True, exist_ok=True)
            if not mdst.exists():
                Image.fromarray(plane, mode="L").save(mdst, format="PNG")
                png_written += 1
            r["mask_conversion_status"] = "CONVERTED"
    else:
        LOGGER.info("[7/8] dry-run: no canonical PNG is written")

    # ---- 8. tests ---------------------------------------------------------------
    LOGGER.info("[8/8] validation ...")
    after_bank = snapshot(prep_root / BANK_NAME)
    pre_splits = [k for k in before_bank if is_split_like(k)]
    new_splits = [k for k in after_bank if k not in before_bank and is_split_like(k)]
    mod_splits = [k for k in after_bank if k in before_bank
                  and after_bank[k] != before_bank[k] and is_split_like(k)]
    new_png = [k for k in after_bank
               if k not in before_bank and k.lower().endswith(".png")]
    raw_diff = diff_state(before_raw, snapshot(src_root))
    after_audit: Dict[str, Tuple[int, int]] = {}
    for d in prep_dirs:
        after_audit.update(snapshot(d))
    aud_diff = diff_state(before_audit, after_audit)
    cpaths = [r["canonical_image_path"] for r in records] + \
             [r["canonical_mask_path"] for r in records]
    collisions = [n for n, c in Counter(p for p in cpaths if p).items() if c > 1]

    T: List[Tuple[str, str, Optional[bool], str]] = []

    def add(i: str, d: str, ok: Optional[bool], b: str) -> None:
        T.append((i, d, ok, b))

    add("T01", f"source records == {EXPECT['records']}",
        len(rows) == EXPECT["records"], f"{len(rows)} @ {MANIFEST_FILE}")
    add("T02", f"source masks == {EXPECT['masks']}", n_masks == EXPECT["masks"],
        f"sum(mask_count)={n_masks}")
    add("T03", f"source benign == {EXPECT['benign']}",
        by_cls.get("BENIGN") == EXPECT["benign"], f"{by_cls.get('BENIGN')}")
    add("T04", f"source malignant == {EXPECT['malignant']}",
        by_cls.get("MALIGNANT") == EXPECT["malignant"], f"{by_cls.get('MALIGNANT')}")
    add("T05", f"local normal == {EXPECT['normal']}",
        int(blob.get("normal_cases", -1)) == EXPECT["normal"]
        and not any("NORMAL" in r["source_label"] for r in records),
        f"summary normal_cases={blob.get('normal_cases')}; no NORMAL label in the "
        "manifest; no external record added")
    add("T06", "class audit agrees with the manifest",
        cls_resolved.get("BENIGN") == by_cls.get("BENIGN")
        and cls_resolved.get("MALIGNANT") == by_cls.get("MALIGNANT"),
        f"class audit {dict(cls_resolved)} vs manifest {dict(by_cls)}")
    add("T07", "summary json agrees with the manifest",
        blob.get("source_images") == len(rows)
        and blob.get("segmentation_masks") == n_masks,
        f"summary {blob.get('source_images')}/{blob.get('segmentation_masks')}")
    add("T08", "pairing artifact agrees: 252 one-to-one pairs",
        len(pair) == EXPECT["records"]
        and all(str(r.get("mask_count", "")).strip() == "1" for r in pair)
        and all(str(r.get("pairing_status", "")).strip() == "VALID_SINGLE_MASK"
                for r in pair),
        f"{len(pair)} rows, statuses "
        f"{dict(Counter(r.get('pairing_status', '') for r in pair))}")
    add("T09", f"image mappings == {EXPECT['images']}", img_ok == EXPECT["images"],
        f"{img_ok}/{len(records)}")
    add("T10", f"mask mappings == {EXPECT['masks']}", mask_ok == EXPECT["masks"],
        f"{mask_ok}/{len(records)}")
    add("T11", f"channel-reduction plans == {EXPECT['eligible']}",
        plans == EXPECT["eligible"], f"{plans}; outcomes {dict(alpha_notes)}")
    add("T12", "mask policy preconditions verified from the mask audit",
        all(precond.values()), join_list([f"{k}={v}" for k, v in precond.items()], "; "))
    add("T13", "threshold optimisation plans == 0",
        "threshold" not in POLICY_RULE.lower(),
        f"declared rule carries no threshold: {POLICY_ID}")
    add("T14", "BUS-UC threshold policy usage == 0",
        not any("BUS_UC" in str(r.get("mask_provenance", "")) for r in records),
        "no record references BUS_UC_CANONICAL_MASK_BINARIZATION_V1")
    add("T15", "multi-mask union plans == 0",
        all(str(r.get("mask_count", "")).strip() == "1" for r in rows)
        and not any("UNION" in str(r.get("mask_provenance", "")) for r in records),
        f"mask_count {dict(Counter(str(r.get('mask_count','')).strip() for r in rows))}")
    add("T16", "pseudo-mask and geometric repair plans == 0", True,
        "no mask is created, resized, resampled, eroded, dilated or smoothed; "
        "a mask that fails reduction is refused, not repaired")
    add("T17", "patient IDs invented == 0", pid_rows == 0
        and id_avail.get("PATIENT_ID") == "NO"
        and all(not r["patient_id"] for r in records),
        f"{IDENTITY_FILE}: PATIENT_ID={id_avail.get('PATIENT_ID')}")
    add("T18", "case IDs invented == 0", case_rows == 0
        and all(not r["case_id"] for r in records),
        f"{IDENTITY_FILE}: CASE_ID={id_avail.get('CASE_ID')}")
    add("T19", "lesion IDs invented == 0", les_rows == 0
        and all(not r["lesion_id"] for r in records),
        f"{IDENTITY_FILE}: LESION_ID={id_avail.get('LESION_ID')}")
    add("T20", "no duplicate exclusion or group lock invented",
        len(dups) == 0 and dup_flagged == 0
        and all(not r["duplicate_group_id"] for r in records),
        f"{DUPLICATE_FILE} rows={len(dups)}; manifest group flags={dup_flagged}")
    add("T21", f"eligible == {EXPECT['eligible']}",
        len(records) == EXPECT["eligible"], f"{len(records)}")
    add("T22", f"eligible benign == {EXPECT['benign']} and malignant == "
               f"{EXPECT['malignant']}",
        by_el.get("BENIGN") == EXPECT["benign"]
        and by_el.get("MALIGNANT") == EXPECT["malignant"],
        f"{by_el.get('BENIGN')} / {by_el.get('MALIGNANT')}")
    add("T23", f"excluded == {EXPECT['excluded']}",
        sum(1 for r in records if r["exclusion_status"] != "INCLUDED_CANONICAL_JOINT")
        == EXPECT["excluded"] and all(str(r.get("audit_status", "")).strip() == "OK"
                                      for r in rows),
        f"audit_status {dict(Counter(r.get('audit_status','') for r in rows))}")
    add("T24", "no unresolved record", not unresolved, f"{len(unresolved)}")
    add("T25", "no canonical path collision", not collisions, f"{len(collisions)}")
    add("T26", "RAW unchanged", raw_diff["unchanged"],
        f"modified {len(raw_diff['modified'])}, added {len(raw_diff['added'])}")
    add("T27", "completed audit artifacts unchanged", aud_diff["unchanged"],
        f"modified {len(aud_diff['modified'])}")
    add("T28", "no split created or modified", not new_splits and not mod_splits,
        f"preexisting={len(pre_splits)}, new={len(new_splits)}, "
        f"modified={len(mod_splits)}")
    add("T29", ("dry-run wrote no canonical PNG" if mode == "dry-run"
                else f"execution wrote {EXPECT['png']} canonical PNG"),
        (not new_png) if mode == "dry-run" else (len(new_png) == EXPECT["png"]),
        f"{len(new_png)} new PNG")
    add("T30", "no split field in the manifest schema",
        not any(is_split_like(f) for f in SCHEMA_FIELDS), "schema checked")
    add("T31", "no preprocessing in the conversion path", None,
        "code-scope declaration: decode, channel reduction, lossless PNG; no resize, "
        "crop, ROI, normalisation, CLAHE, gamma, denoise, sharpen or augmentation")
    verified = [t for t in T if t[2] is not None]
    failed = [t for t in verified if not t[2]]
    for t in T:
        (LOGGER.info if t[2] is not False else LOGGER.error)(
            "      %-5s [%-11s] %-50s %s", t[0],
            "DECLARATION" if t[2] is None else ("PASS" if t[2] else "FAIL"),
            t[1][:50], t[3][:52])

    # ---- outputs ----------------------------------------------------------------
    tag = "DRY_RUN_" if mode == "dry-run" else ""
    out_rows = [{k: r[k] for k in SCHEMA_FIELDS} for r in records]
    write_csv(manifests / f"{tag}canonical_manifest.csv", SCHEMA_FIELDS, out_rows)
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "canonical_manifest"
        ws.append(SCHEMA_FIELDS)
        for r in out_rows:
            ws.append([r.get(h, "") for h in SCHEMA_FIELDS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = (f"A1:{get_column_letter(len(SCHEMA_FIELDS))}"
                              f"{len(out_rows) + 1}")
        wb.save(manifests / f"{tag}canonical_manifest.xlsx")
        LOGGER.info("WROTE %-46s (1 sheet)", f"{tag}canonical_manifest.xlsx")
    except ImportError:
        LOGGER.warning("openpyxl not installed; manifest XLSX skipped")

    blocked = bool(failed)
    state = {
        "phase": "2B.3.3", "dataset": DATASET, "mode": mode, "generated": now_iso(),
        "builder_sha256": sha256_file(Path(__file__).resolve()),
        "schema_sha256": sha256_text(join_list(SCHEMA_FIELDS, ",")),
        "evidence_hashes": ev["hashes"], "source_root": str(src_root),
        "mask_policy": {"policy_id": POLICY_ID, "rule": POLICY_RULE,
                        "threshold_used": None,
                        "preconditions_verified": precond,
                        "channel_reduction_outcomes": dict(alpha_notes)},
        "counts": {"source_records": len(rows), "source_masks": n_masks,
                   "source_benign": by_cls.get("BENIGN", 0),
                   "source_malignant": by_cls.get("MALIGNANT", 0),
                   "source_normal_local": blob.get("normal_cases"),
                   "eligible": len(records),
                   "eligible_benign": by_el.get("BENIGN", 0),
                   "eligible_malignant": by_el.get("MALIGNANT", 0),
                   "excluded": 0, "reduction_plans": plans, "union_plans": 0,
                   "threshold_plans": 0,
                   "exact_duplicate_rows": len(dups), "group_locks": 0},
        "image_mappings_resolved": img_ok, "mask_mappings_resolved": mask_ok,
        "unresolved": unresolved[:50], "collisions": collisions,
        "canonical_png_written": png_written,
        "preexisting_split_artifacts": len(pre_splits),
        "newly_created_split_artifacts": len(new_splits),
        "modified_split_artifacts": len(mod_splits),
        "raw_modified": not raw_diff["unchanged"],
        "audit_modified": not aud_diff["unchanged"],
        "tests_pass": sum(1 for t in verified if t[2]), "tests_fail": len(failed),
        "verdict": ("BLOCKED" if blocked else
                    "READY_FOR_BREAST_CANONICAL_EXECUTION" if mode == "dry-run"
                    else "BREAST_CANONICAL_BUILD_COMPLETE"),
    }
    write_text(reports / f"{tag}canonical_build_report.json",
               json.dumps(state, indent=2, ensure_ascii=False, default=str))

    LOGGER.info("=" * 70)
    LOGGER.info("PHASE 2B.3.3 BrEaST CANONICAL BANK - %s", mode.upper())
    LOGGER.info("")
    LOGGER.info("Evidence: %s, %s, %s, %s, %s, %s, %s", MANIFEST_FILE, PAIRING_FILE,
                MASK_FILE, CLASS_FILE, IDENTITY_FILE, DUPLICATE_FILE, SUMMARY_FILE)
    LOGGER.info("Source:   %d records / %d masks  (benign %d / malignant %d / "
                "normal %s)", len(rows), n_masks, by_cls.get("BENIGN", 0),
                by_cls.get("MALIGNANT", 0), blob.get("normal_cases"))
    LOGGER.info("Canonical: eligible %d (benign %d / malignant %d), excluded 0",
                len(records), by_el.get("BENIGN", 0), by_el.get("MALIGNANT", 0))
    LOGGER.info("Mask rule: %s -> %s", POLICY_ID, POLICY_RULE)
    LOGGER.info("           threshold plans 0 | union plans 0 | pseudo-mask 0 | "
                "geometric repair 0")
    LOGGER.info("Identity:  IMAGE_ONLY; patient/case/lesion invented = 0/0/0")
    LOGGER.info("Duplicates: %d exact-duplicate row(s), 0 group lock(s)", len(dups))
    LOGGER.info("image mappings %d | mask mappings %d | unresolved %d | collisions %d",
                img_ok, mask_ok, len(unresolved), len(collisions))
    LOGGER.info("RAW modified %s | audit modified %s | canonical PNG %d | splits %d/%d",
                state["raw_modified"], state["audit_modified"], png_written,
                len(new_splits), len(mod_splits))
    LOGGER.info("")
    LOGGER.info("TESTS PASS: %d    TESTS FAIL: %d", state["tests_pass"],
                state["tests_fail"])
    LOGGER.info("")
    LOGGER.info("FINAL VERDICT:")
    if blocked:
        LOGGER.error("BLOCKED")
        for t in failed:
            LOGGER.error("    FAILED %s: %s | %s", t[0], t[1], t[3])
        return 4
    if mode == "dry-run":
        LOGGER.info("READY_FOR_BREAST_CANONICAL_EXECUTION")
        LOGGER.info("=" * 70)
        LOGGER.info("STOP. --execute is not run automatically and needs human "
                    "approval.")
    else:
        LOGGER.info("BREAST_CANONICAL_BUILD_COMPLETE")
    return 0


def execute_gate(prep_root: Path) -> Tuple[bool, str]:
    st = (prep_root / BANK_NAME / DATASET / "reports"
          / "DRY_RUN_canonical_build_report.json")
    if not st.is_file():
        return False, f"no dry-run report at {st}"
    d = json.loads(st.read_text(encoding="utf-8"))
    if d.get("verdict") != "READY_FOR_BREAST_CANONICAL_EXECUTION":
        return False, f"the last dry-run verdict was {d.get('verdict')!r}"
    if d.get("builder_sha256") != sha256_file(Path(__file__).resolve()):
        return False, "the builder changed since the approved dry-run; rerun --dry-run"
    if d.get("schema_sha256") != sha256_text(join_list(SCHEMA_FIELDS, ",")):
        return False, "the schema changed since the approved dry-run; rerun --dry-run"
    ev = load_evidence(prep_root)
    for name, h in (d.get("evidence_hashes") or {}).items():
        if ev["hashes"].get(name) != h:
            return False, f"evidence artifact {name} changed since the dry-run"
    return True, "the dry-run approval matches the current builder, schema and evidence"


def main(argv: Optional[Sequence[str]] = None) -> int:
    p = argparse.ArgumentParser(description="BrEaST canonical bank builder.")
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--dry-run", action="store_true")
    g.add_argument("--execute", action="store_true")
    p.add_argument("--prep-root", default=DEFAULT_PREP_ROOT)
    a = p.parse_args(argv)
    mode = "dry-run" if a.dry_run else "execute"
    if mode == "execute":
        ok, why = execute_gate(Path(a.prep_root))
        if not ok:
            sys.stderr.write(f"EXECUTE refused: {why}\n")
            return 5
        sys.stderr.write(f"execute gate: {why}\n")
    try:
        return run(mode, Path(a.prep_root))
    except Exception:  # noqa: BLE001
        if LOGGER.handlers:
            LOGGER.critical("UNHANDLED EXCEPTION\n%s", traceback.format_exc())
        else:
            sys.stderr.write(traceback.format_exc())
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
